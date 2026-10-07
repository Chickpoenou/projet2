"""Construit le classeur SUIVI_GC5 (.xlsm) : cahier de texte, présences, exposés,
rapports, fiches étudiants, suivi des enseignants et exports PDF.

Usage :
    python build_classeur.py SORTIE.xlsm [--etudiants SOURCE.xlsm] [--demo]

--etudiants : classeur d'origine dont la feuille « Liste » fournit les étudiants
              (N°, Nom et prénoms, Groupe, Email, Contact). Sans cette option, la
              liste est vide (version publiable sans données personnelles).
--demo      : ajoute des séances et présences fictives (tests uniquement).
"""

import argparse
import datetime as dt
import os
import random
import sys

import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell as rc, xl_col_to_name

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from vba_project import build_vba_project  # noqa: E402

ICI = os.path.dirname(os.path.abspath(__file__))
VBA_DIR = os.path.join(ICI, "..", "vba")
MDP = "GC5BTP"

# ------------------------------------------------------------------ palette
MARINE = "#1F3864"
BLEU = "#2E75B6"
BLEU_CLAIR = "#DDEBF7"
PALE = "#F3F7FC"
GRIS_TXT = "#595959"
BORDURE = "#BFBFBF"
ORANGE = "#ED7D31"
VERT = "#70AD47"
ROUGE = "#C00000"
OR = "#BF9000"
BLANC = "#FFFFFF"
STATUTS = {  # fond, texte
    "P": ("#C6EFCE", "#006100"),
    "A": ("#FFC7CE", "#9C0006"),
    "R": ("#FFEB9C", "#9C5700"),
    "E": ("#D9D9D9", "#404040"),
}
POLICE = "Segoe UI"

# ------------------------------------------------------------- référentiels
MATIERES = [
    # code, UE, intitulé, abrégé, enseignant(s), cours, TD, TP, couleur
    ("01GEC2301", "GEC2301", "Droits des Travaux publics (Marchés et Passation de marchés)", "Droit TP",
     "GIBIGAYE / SEKLOKA", 42, 0, 0, "#FCE4D6"),
    ("02GEC2301", "GEC2301", "CAO-DAO Appliqué", "CAO-DAO", "GODONOU", 35, 21, 0, "#DDEBF7"),
    ("01GEC2302", "GEC2302", "Sortie pédagogique", "Sortie péda.", "HOUANOU K. A.", 0, 0, 42, "#E2EFDA"),
    ("02GEC2302", "GEC2302", "Travaux Pratiques Spécialisés", "TP spécialisés",
     "GODONOU / SOHOUNHLOUE", 11, 45, 0, "#FFF2CC"),
    ("GEC2303", "GEC2303", "Introduction aux Eurocodes", "Eurocodes", "HOUANOU K. A.", 35, 21, 0, "#EADCF8"),
    ("GEC2304", "GEC2304", "Conception et Calcul de Ponts", "Ponts", "DOKO V.", 56, 0, 0, "#D0E8F2"),
    ("GEC2305", "GEC2305", "Projet de Construction (BA et CM)", "Projet BA-CM", "GIBIGAYE", 56, 0, 0, "#F8D7DA"),
    ("GEC2306", "GEC2306", "Métré et Estimation de Prix", "Métré", "MILOHIN", 42, 0, 0, "#E8E3D3"),
    ("01GEC2307", "GEC2307", "Entrepreneuriat et Leadership", "Entrepreneuriat", "ALE", 28, 0, 0, "#D5F5E3"),
    ("02GEC2307", "GEC2307", "Législation du travail", "Législation", "Prof X", 28, 0, 0, "#FDEBD0"),
    ("GEC2308", "GEC2308", "Management des Projets", "Management", "ASSOGBA Gauthier", 28, 0, 0, "#E5E7E9"),
]
COULEUR = {m[0]: m[8] for m in MATIERES}
ABREGE = {m[0]: m[3] for m in MATIERES}
ENSEIGNANT = {m[0]: m[4] for m in MATIERES}

ENSEIGNANTS = ["ALE", "ASSOGBA Gauthier", "DOKO V.", "GIBIGAYE", "GODONOU", "HOUANOU K. A.", "MILOHIN",
               "Prof X", "SEKLOKA", "SOHOUNHLOUE"]

JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]
EDT = [
    # jour, début, fin, code, salle
    ("Lundi", 8, 12, "GEC2304", "B 2650"),
    ("Lundi", 14, 17, "01GEC2301", "B 2650"),
    ("Mardi", 8, 12, "GEC2306", "B 2650"),
    ("Mardi", 14, 18, "GEC2303", "B-2610"),
    ("Mercredi", 8, 11, "GEC2305", "B 2650"),
    ("Mercredi", 11, 13, "02GEC2307", "B 2650"),
    ("Jeudi", 8, 11, "02GEC2302", "B 2650"),
    ("Jeudi", 11, 13, "01GEC2307", "B 2650"),
    ("Jeudi", 14, 18, "02GEC2301", "B-2610"),
    ("Vendredi", 14, 16, "GEC2308", ""),
    ("Vendredi", 16, 18, "01GEC2302", ""),
    ("Samedi", 8, 12, "GEC2304", "B 2650"),
]

PARAMS = [
    ("pEcole", "Établissement", "École Polytechnique d'Abomey-Calavi"),
    ("pDepartement", "Département", "Département de Génie Civil"),
    ("pOption", "Option", "BTP"),
    ("pNiveau", "Niveau / promotion", "GC5"),
    ("pAnnee", "Année académique", "2026-2027"),
    ("pSemestre", "Semestre", "Semestre 9"),
    ("pDebut", "Début de la période", dt.date(2026, 9, 28)),
    ("pFin", "Fin de la période", dt.date(2027, 1, 22)),
    ("pSeuil", "Seuil d'alerte : absences par matière (rouge au-delà)", 3),
    ("pMarge", "Marge avant le début d'un cours (minutes)", 30),
    ("pDelegue", "Délégué de classe (nom affiché sur les PDF)", ""),
]

LISTES = {
    "LstStatuts": ("Statut de présence", ["P", "A", "R", "E"]),
    "LstProf": ("Présence du professeur", ["Présent", "Absent", "Retard", "Remplacé"]),
    "LstTypes": ("Type de séance", ["Cours", "TD", "TP", "Exposé", "Évaluation", "Sortie", "Rattrapage"]),
    "LstStatutExpose": ("Statut d'exposé", ["À venir", "Présenté", "Reporté", "Annulé"]),
    "LstTypeRapport": ("Type de rapport", ["Individuel", "Groupe"]),
    "LstJours": ("Jours", JOURS),
    "LstModeExport": ("Mode d'export", ["Un PDF par groupe", "Un seul PDF (tous les groupes)"]),
}
SIGNIF = {"P": "Présent", "A": "Absent", "R": "Retard", "E": "Excusé"}

NB_LIGNES_TABLE = 3000      # portée des validations / mises en forme sous les tableaux


def _args(txt, debut):
    """Découpe les arguments d'un appel de fonction à partir de la parenthèse ouvrante."""
    prof, crochet, args, cur, i = 0, 0, [], "", debut + 1
    while i < len(txt):
        ch = txt[i]
        if ch == '"':
            j = txt.index('"', i + 1)
            cur += txt[i:j + 1]
            i = j + 1
            continue
        if ch == "[":
            crochet += 1
        elif ch == "]":
            crochet -= 1
        elif ch == "(":
            prof += 1
        elif ch == ")":
            if prof == 0 and crochet == 0:
                args.append(cur)
                return args, i
            prof -= 1
        elif ch == "," and prof == 0 and crochet == 0:
            args.append(cur)
            cur = ""
            i += 1
            continue
        cur += ch
        i += 1
    raise ValueError(txt)


def sans_xlookup(f):
    """XLOOKUP(v, clés, retour[, défaut]) -> IFERROR(INDEX(retour,MATCH(v,clés,0)),défaut)."""
    while "XLOOKUP(" in f:
        d = f.index("XLOOKUP(")
        args, fin = _args(f, d + 7)
        v, cles, ret = args[0], args[1], args[2]
        defaut = args[3] if len(args) > 3 else '""'
        f = f[:d] + f"IFERROR(INDEX({ret},MATCH({v},{cles},0)),{defaut})" + f[fin + 1:]
    return f


def _compat(methode):
    """Enveloppe une méthode d'écriture de formule : XLOOKUP -> INDEX/MATCH (Excel 2016/2019)."""
    def f(self, *args, **kw):
        args = [sans_xlookup(a) if isinstance(a, str) and "XLOOKUP(" in a else a for a in args]
        return methode(self, *args, **kw)
    return f


from xlsxwriter.worksheet import Worksheet as _WS  # noqa: E402
for _m in ("write_formula", "write_array_formula", "write_dynamic_array_formula"):
    setattr(_WS, _m, _compat(getattr(_WS, _m)))


def heure(h, m=0):
    return dt.time(h, m)


def lire_etudiants(chemin):
    import openpyxl
    wb = openpyxl.load_workbook(chemin, data_only=True)
    ws = wb["Liste"]
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[2]:
            nom = str(r[2]).strip()
            role = "Animateur" if nom in ("BIO SEKO Aïmane",) else ""
            contact = r[5]
            out.append([r[1], nom, r[3], role, r[4] or "", contact if contact else ""])
    return out


TABLES_REPRISES = ["tblSeances", "tblPresences", "tblExposes", "tblRapports", "tblRemises", "tblGroupes",
                   "tblEtudiants", "tblListeGroupes", "tblEnseignants"]


def lire_reprise(chemin):
    """Lit les saisies d'un ancien classeur : {table: [ {en-tête: valeur} ]}, {paramètre: valeur}."""
    import openpyxl
    wb = openpyxl.load_workbook(chemin)
    tables, params = {}, {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            if t.name not in TABLES_REPRISES:
                continue
            lignes = list(ws[t.ref])
            entetes = [c.value for c in lignes[0]]
            out = []
            for ligne in lignes[1:]:
                d = {}
                for h, c in zip(entetes, ligne):
                    v = c.value
                    if v is None or not isinstance(v, (str, int, float, dt.date, dt.time, dt.datetime)):
                        continue
                    if isinstance(v, str) and v.startswith("="):
                        continue
                    d[h] = v
                if d:
                    out.append(d)
            if out:
                tables[t.name] = out
    for nom in wb.defined_names:
        if nom.startswith("p"):
            try:
                dest = list(wb.defined_names[nom].destinations)[0]
                params[nom] = wb[dest[0]][dest[1].replace("$", "")].value
            except Exception:
                pass
    return tables, params


# ======================================================================
class Classeur:
    def __init__(self, sortie, etudiants, demo, apercu=False):
        self.apercu = apercu
        self.wb = xlsxwriter.Workbook(sortie, {"use_future_functions": True})
        self.etudiants = etudiants
        self.demo = demo
        self._fmts = {}
        self.wb.set_vba_name("ThisWorkbook")
        self.wb.set_properties({"title": "Suivi GC5 - cahier de texte et présences",
                                "author": "Délégué GC5", "company": "EPAC"})
        self.wb.set_size(1600, 1000)

    # -------------------------------------------------------------- formats
    def f(self, **kw):
        base = {"font_name": POLICE, "font_size": 10, "valign": "vcenter"}
        base.update(kw)
        cle = tuple(sorted(base.items()))
        if cle not in self._fmts:
            self._fmts[cle] = self.wb.add_format(base)
        return self._fmts[cle]

    def saisie(self, **kw):
        kw.setdefault("bg_color", "#FFFDF2")
        kw.setdefault("border", 1)
        kw.setdefault("border_color", "#F4B183")
        return self.f(locked=False, **kw)

    # ------------------------------------------------------------ éléments
    def feuille(self, nom, code, onglet, titre, sous_titre, larg, dernier_col):
        ws = self.wb.add_worksheet(nom)
        ws.set_vba_name(code)
        ws.set_tab_color(onglet)
        ws.hide_gridlines(2)
        ws.set_zoom(90)
        ws.set_column(0, 0, 2)
        for i, w in enumerate(larg):
            ws.set_column(i + 1, i + 1, w)
        ws.set_row(0, 8)
        ws.set_row(1, 34)
        ws.merge_range(1, 1, 1, dernier_col, titre,
                       self.f(bold=True, font_size=17, font_color=BLANC, bg_color=MARINE, indent=1))
        ws.merge_range(2, 1, 2, dernier_col, sous_titre,
                       self.f(italic=True, font_color=GRIS_TXT, font_size=9, indent=1, text_wrap=True))
        ws.set_row(2, 26)
        ws.write_url(0, 1, "internal:'ACCUEIL'!A1", self.f(font_size=7, font_color=BLEU, underline=1),
                     string="⌂ Accueil")
        ws.write_url(0, 2, "internal:'MODE D''EMPLOI'!A1", self.f(font_size=7, font_color=BLEU, underline=1),
                     string="?  Mode d'emploi")
        if self.apercu:
            ws.set_landscape()
            ws.fit_to_pages(1, 1)
            ws.print_area(0, 0, 47, max(dernier_col, 14))
        ws.protect(MDP, {"autofilter": True, "format_columns": True, "format_rows": True,
                         "select_locked_cells": True, "select_unlocked_cells": True})
        return ws

    def bouton(self, ws, ligne, col, texte, cible=None, macro=None, couleur=BLEU, larg=170, haut=34,
               dx=0, dy=0, taille=10):
        opts = {
            "width": larg, "height": haut, "x_offset": dx, "y_offset": dy,
            "font": {"name": POLICE, "bold": True, "color": BLANC, "size": taille},
            "align": {"vertical": "middle", "horizontal": "center"},
            "fill": {"color": couleur},
            "line": {"color": couleur, "width": 1},
            "object_position": 3,
        }
        if cible:
            opts["url"] = "internal:'" + cible.replace("'", "''") + "'!A1"
            opts["tip"] = f"Aller à {cible}"
        if macro:
            opts["description"] = f"macro:{macro}"
        ws.insert_textbox(ligne, col, texte, opts)

    def rangee(self, ws, ligne, items, haut=28, dy=1, x0=0, gap=12):
        """Aligne des boutons sur une ligne, à partir de la colonne B."""
        x = x0
        for texte, kw in items:
            larg = kw.pop("larg", 150)
            self.bouton(ws, ligne, 1, texte, larg=larg, haut=haut, dx=x, dy=dy, **kw)
            x += larg + gap

    def section(self, ws, ligne, c1, c2, texte, couleur=BLEU):
        if c1 == c2:
            ws.write(ligne, c1, texte, self.f(bold=True, font_color=BLANC, bg_color=couleur, indent=1))
        else:
            ws.merge_range(ligne, c1, ligne, c2, texte,
                           self.f(bold=True, font_color=BLANC, bg_color=couleur, indent=1))

    def table(self, ws, ligne, col, nom, colonnes, donnees=None, style="Table Style Medium 2", total=False):
        """colonnes : liste de dicts {header, format, formula, width?, total_function?}."""
        import re
        for c in colonnes:
            if "formula" in c:
                fml = re.sub(r"\[@([^\[\]]+)\]", r"[@[\1]]", c["formula"])
                c["formula"] = re.sub(r"(?<![\w\]])\[@", nom + "[@", fml)
                c["formula"] = sans_xlookup(c["formula"])
        reprise = getattr(self, "reprise", {}).get(nom)
        if reprise:
            donnees = [[None if "formula" in c else ligne.get(c["header"]) for c in colonnes] for ligne in reprise]
        n = max(len(donnees or []), 1)
        derniere = ligne + n + (1 if total else 0)
        opts = {"name": nom, "style": style, "columns": colonnes, "autofilter": True}
        if donnees:
            opts["data"] = donnees
        if total:
            opts["total_row"] = True
        ws.add_table(ligne, col, derniere, col + len(colonnes) - 1, opts)
        ws.set_row(ligne, 30)
        return derniere

    def cf_statuts(self, ws, plage):
        for s, (fond, txt) in STATUTS.items():
            ws.conditional_format(plage, {"type": "cell", "criteria": "==", "value": f'"{s}"',
                                          "format": self.f(bg_color=fond, font_color=txt, bold=True)})

    def dv_liste(self, ws, plage, source, message=None, libre=False):
        opts = {"validate": "list", "source": source, "ignore_blank": True}
        if libre:  # liste proposée, saisie libre acceptée
            opts["error_type"] = "information"
        if message:
            opts.update({"input_title": message[0], "input_message": message[1]})
        ws.data_validation(plage, opts)

    def let(self, formule):
        """Préfixe les variables LET (notées §x) par _xlpm. comme l'exige le format de fichier."""
        import re
        return re.sub(r"§(\w+)", r"_xlpm.\1", formule)

    # ===================================================================
    def construire(self):
        self.accueil()
        self.guide()
        self.cours_du_jour()
        self.cahier()
        self.presences()
        self.exposes()
        self.rapports()
        self.remises()
        self.fiche()
        self.suivi()
        self.recap()
        self.export()
        self.etudiants_f()
        self.groupes()
        self.matieres()
        self.edt()
        self.parametres()
        self.impression()
        self.noms()
        self.vba()
        self.wb.worksheets()[0].activate()
        self.wb.close()

    # ---------------------------------------------------------- ACCUEIL
    def accueil(self):
        ws = self.feuille("ACCUEIL", "shAccueil", MARINE,
                          "CAHIER DE TEXTE  ·  PRÉSENCES  ·  SUIVI DES ÉTUDIANTS ET DES ENSEIGNANTS",
                          "", [11] * 12 + [3, 3], 12)
        self.ws_accueil = ws
        ws.write_formula("B3", '=pEcole&"  —  "&pDepartement&"  —  "&pNiveau&" "&pOption&"  —  Année académique "&pAnnee&"  —  "&pSemestre',
                         self.f(italic=True, font_color=GRIS_TXT, font_size=10, indent=1))
        ws.set_row(3, 22)
        ws.merge_range("B4:H4", "", self.f())
        ws.write_formula("B4", "=TODAY()", self.f(num_format="[$-40C]dddd d mmmm yyyy", bold=True,
                                                    font_size=12, font_color=MARINE, indent=1))
        ws.merge_range("I4:M4", "", self.f())
        ws.write_formula("I4", '="Mis à jour à "&TEXT(NOW(),"hh:mm")&"  (F9 pour actualiser)"',
                         self.f(align="right", font_color=GRIS_TXT, font_size=9))
        # Cellules de calcul (colonne P masquée)
        ws.set_column("P:P", None, None, {"hidden": True})
        jour = "WEEKDAY(TODAY(),2)"
        maint = "MOD(NOW(),1)"
        ws.write_dynamic_array_formula(
            "P6:P6",
            f"=IFERROR(MATCH(1,(tblEDT[N° jour]={jour})*(tblEDT[Début]-pMarge/1440<={maint})*(tblEDT[Fin]>{maint}),0),0)")
        ws.write_array_formula("P7:P7", '{=IFERROR(MATCH(MIN(IF(ISNUMBER(tblEDT[Minutes avant]),IF(tblEDT[Minutes avant]>pMarge,tblEDT[Minutes avant]))),tblEDT[Minutes avant],0),0)}')
        ws.write_formula("P8", "=IF(P7=0,\"\",INDEX(tblEDT[Minutes avant],P7))")

        # Cartes cours du moment / prochain cours
        carte_t = self.f(bold=True, font_color=BLANC, bg_color=BLEU, indent=1)
        carte_v = self.f(bold=True, font_size=15, font_color=MARINE, bg_color=PALE, indent=1)
        carte_s = self.f(font_color=GRIS_TXT, bg_color=PALE, indent=1)
        carte_ok = self.f(bold=True, font_color="#006100", bg_color=PALE, indent=1)
        ws.merge_range("B6:G6", "●  COURS DU MOMENT", carte_t)
        ws.merge_range("I6:M6", "➜  PROCHAIN COURS", self.f(bold=True, font_color=BLANC, bg_color=OR, indent=1))
        for r in (6, 7, 8):
            ws.merge_range(r, 1, r, 6, "", carte_s)
            ws.merge_range(r, 8, r, 12, "", carte_s)
        ws.set_row(6, 30)
        ws.write_formula("B7", '=IF($P$6=0,"Aucun cours en ce moment",INDEX(tblEDT[Code matière],$P$6)&"  ·  "&INDEX(tblEDT[Matière],$P$6))', carte_v)
        ws.write_formula("B8", '=IF($P$6=0,"Profitez-en pour mettre à jour les rapports et exposés.",TEXT(INDEX(tblEDT[Début],$P$6),"hh:mm")&" – "&TEXT(INDEX(tblEDT[Fin],$P$6),"hh:mm")&"   ·   "&INDEX(tblEDT[Enseignant],$P$6)&"   ·   Salle "&INDEX(tblEDT[Salle],$P$6))', carte_s)
        ws.write_formula("B9", '=IF($P$6=0,"",IF(COUNTIFS(tblSeances[Date],TODAY(),tblSeances[Code matière],INDEX(tblEDT[Code matière],$P$6))>0,"✔ Séance enregistrée dans le cahier de texte","⚠ Séance pas encore démarrée : cliquez sur « DÉMARRER LE COURS »"))', carte_ok)
        ws.conditional_format("B9", {"type": "formula", "criteria": '=LEFT($B$9,1)="⚠"',
                                     "format": self.f(font_color="#C55A11")})
        ws.set_row(7, 20)
        ws.write_formula("I7", '=IF($P$7=0,"—",INDEX(tblEDT[Code matière],$P$7)&"  ·  "&INDEX(tblEDT[Matière],$P$7))', carte_v)
        ws.write_formula("I8", '=IF($P$7=0,"",INDEX(tblEDT[Jour],$P$7)&" "&TEXT(INDEX(tblEDT[Début],$P$7),"hh:mm")&"   ·   "&INDEX(tblEDT[Enseignant],$P$7)&"   ·   Salle "&INDEX(tblEDT[Salle],$P$7))', carte_s)
        ws.write_formula("I9", '=IF($P$7=0,"","Dans "&IF($P$8>=1440,INT($P$8/1440)&" j ","")&INT(MOD($P$8,1440)/60)&" h "&TEXT(MOD($P$8,60),"00")&" min")', carte_ok)

        # Indicateurs
        kpis = [
            ("SÉANCES ENREGISTRÉES", '=COUNTIF(tblSeances[ID],"S*")',
             '="dont "&COUNTIF(tblSeances[Présence prof],"Absent")&" sans professeur"', BLEU, "0"),
            ("HEURES EFFECTUÉES", '=SUMIFS(tblSeances[Durée (h)],tblSeances[Présence prof],"<>Absent")',
             '="sur "&SUM(tblMatieres[Volume total (h)])&" h prévues ("&TEXT(IFERROR(SUMIFS(tblSeances[Durée (h)],tblSeances[Présence prof],"<>Absent")/SUM(tblMatieres[Volume total (h)]),0),"0%")&")"',
             VERT, '0.0" h"'),
            ("TAUX DE PRÉSENCE", '=IFERROR((COUNTIF(tblPresences[Statut],"P")+COUNTIF(tblPresences[Statut],"R"))/COUNTIF(tblPresences[Statut],"?"),"—")',
             '=COUNTIF(tblPresences[Statut],"A")&" absences · "&COUNTIF(tblPresences[Statut],"R")&" retards"', MARINE, "0%"),
            ("ÉTUDIANTS EN ALERTE", "=COUNTIF('RÉCAP ABSENCES'!$T$7:$T$206,\">\"&pSeuil)",
             '="plus de "&pSeuil&" absences dans une matière"', ROUGE, "0"),
            ("RAPPORTS NON REMIS", '=COUNTIF(tblRemises[Statut],"Non remis")',
             '=COUNTIF(tblRemises[Statut],"En retard")&" remis en retard"', ORANGE, "0"),
            ("EXPOSÉS RESTANTS", '=COUNTIF(tblExposes[Statut],"À venir")+COUNTIF(tblExposes[Statut],"Reporté")',
             '=COUNTIF(tblExposes[Statut],"Présenté")&" déjà présentés"', OR, "0"),
        ]
        ws.set_row(11, 36)
        for i, (lib, val, sous, coul, nf) in enumerate(kpis):
            c = 1 + 2 * i
            ws.merge_range(10, c, 10, c + 1, lib, self.f(bold=True, font_size=8, font_color=BLANC, bg_color=coul,
                                                         align="center"))
            ws.merge_range(11, c, 11, c + 1, "", self.f())
            ws.write_formula(11, c, val, self.f(bold=True, font_size=20, font_color=coul, bg_color=PALE,
                                                align="center", num_format=nf))
            ws.merge_range(12, c, 12, c + 1, "", self.f())
            ws.write_formula(12, c, sous, self.f(font_size=8, font_color=GRIS_TXT, bg_color=PALE, align="center",
                                                 text_wrap=True))
        ws.set_row(12, 26)

        # Cours d'aujourd'hui (8 lignes, tri par heure)
        self.section(ws, 14, 1, 12, "COURS D'AUJOURD'HUI")
        entetes = [("B16:C16", "Horaire"), ("D16:D16", "Code"), ("E16:H16", "Matière"),
                   ("I16:J16", "Enseignant"), ("K16:K16", "Salle"), ("L16:M16", "État")]
        th = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1)
        for plage, t in entetes:
            if ":" in plage and plage.split(":")[0] != plage.split(":")[1]:
                ws.merge_range(plage, t, th)
            else:
                ws.write(plage.split(":")[0], t, th)
        cle = "tblEDT[Début]+ROW(tblEDT[Début])/100000"
        for k in range(8):
            r = 16 + k
            fond = PALE if k % 2 else BLANC
            cell = self.f(bg_color=fond, align="center")
            celg = self.f(bg_color=fond, indent=1)
            ws.write_dynamic_array_formula(
                r, 15, r, 15,
                f"=IFERROR(MATCH(SMALL(IF(tblEDT[N° jour]={jour},{cle}),{k + 1}),{cle},0),0)")
            idx = f"$P${r + 1}"
            ws.merge_range(r, 1, r, 2, "", cell)
            ws.write_formula(r, 1, f'=IF({idx}=0,"",TEXT(INDEX(tblEDT[Début],{idx}),"hh:mm")&" – "&TEXT(INDEX(tblEDT[Fin],{idx}),"hh:mm"))', cell)
            ws.write_formula(r, 3, f'=IF({idx}=0,"",INDEX(tblEDT[Code matière],{idx}))', self.f(bg_color=fond, align="center", bold=True))
            ws.merge_range(r, 4, r, 7, "", celg)
            ws.write_formula(r, 4, f'=IF({idx}=0,IF({k}=0,"Pas de cours prévu aujourd\'hui",""),INDEX(tblEDT[Matière],{idx}))', celg)
            ws.merge_range(r, 8, r, 9, "", celg)
            ws.write_formula(r, 8, f'=IF({idx}=0,"",INDEX(tblEDT[Enseignant],{idx}))', celg)
            ws.write_formula(r, 10, f'=IF({idx}=0,"",INDEX(tblEDT[Salle],{idx}))', cell)
            ws.merge_range(r, 11, r, 12, "", cell)
            ws.write_formula(
                r, 11,
                f'=IF({idx}=0,"",IF(COUNTIFS(tblSeances[Date],TODAY(),tblSeances[Code matière],INDEX(tblEDT[Code matière],{idx}))>0,"✔ Enregistré",IF(INDEX(tblEDT[Fin],{idx})<MOD(NOW(),1),"✖ Non saisi",IF(INDEX(tblEDT[Début],{idx})<=MOD(NOW(),1),"● En cours","À venir"))))',
                cell)
        for t, coul in (("✔", "#006100"), ("✖", ROUGE), ("●", "#C55A11")):
            ws.conditional_format("L17:L24", {"type": "formula", "criteria": f'=LEFT($L17,1)="{t}"',
                                              "format": self.f(bold=True, font_color=coul)})

        # Accès rapide
        self.section(ws, 26, 1, 12, "ACCÈS RAPIDE")
        tuiles = [
            ("▶  DÉMARRER LE COURS", None, "DemarrerCours", ORANGE),
            ("✎  Cours du jour / Appel", "COURS DU JOUR", None, ORANGE),
            ("📖  Cahier de texte", "CAHIER DE TEXTE", None, ORANGE),
            ("✔  Présences", "PRÉSENCES", None, ORANGE),
            ("🎤  Exposés", "EXPOSÉS", None, "#C55A11"),
            ("📄  Rapports", "RAPPORTS", None, "#C55A11"),
            ("📥  Remises", "REMISES", None, "#C55A11"),
            ("🔍  Fiche étudiant", "FICHE ÉTUDIANT", None, VERT),
            ("👤  Suivi enseignants", "SUIVI ENSEIGNANTS", None, VERT),
            ("⚠  Récap absences", "RÉCAP ABSENCES", None, VERT),
            ("🖨  Export PDF", "EXPORT PDF", None, VERT),
            ("👥  Étudiants", "ÉTUDIANTS", None, "#7F7F7F"),
            ("🧩  Groupes", "GROUPES", None, "#7F7F7F"),
            ("📚  Matières", "MATIÈRES", None, "#7F7F7F"),
            ("🗓  Emploi du temps", "EMPLOI DU TEMPS", None, "#7F7F7F"),
            ("⚙  Paramètres", "PARAMÈTRES", None, "#7F7F7F"),
            ("💾  Sauvegarder (copie datée)", None, "Sauvegarder", MARINE),
            ("🔒  Mode administrateur", None, "ModeAdministrateur", MARINE),
            ("↻  Actualiser", None, "Actualiser", MARINE),
            ("❓  MODE D'EMPLOI", "MODE D'EMPLOI", None, OR),
        ]
        for i, (txt, cible, macro, coul) in enumerate(tuiles):
            ligne, colonne = 28 + (i // 4) * 3, 1 + (i % 4) * 2
            self.bouton(ws, ligne, colonne, txt, cible=cible, macro=macro, couleur=coul, larg=150, haut=40,
                        dx=2, taille=10)
        # Graphique d'avancement
        ch = self.wb.add_chart({"type": "bar"})
        ch.add_series({
            "name": "Avancement",
            "categories": "='SUIVI ENSEIGNANTS'!$C$7:$C$17",
            "values": "='SUIVI ENSEIGNANTS'!$K$7:$K$17",
            "fill": {"color": BLEU}, "border": {"none": True}, "gap": 60,
            "data_labels": {"value": True, "num_format": "0%", "font": {"size": 8, "color": GRIS_TXT}},
        })
        ch.set_title({"name": "Avancement des matières (heures faites / prévues)",
                      "name_font": {"size": 10, "bold": True, "color": MARINE}})
        ch.set_x_axis({"min": 0, "max": 1, "num_format": "0%", "major_gridlines": {"visible": True,
                       "line": {"color": "#E7E6E6"}}, "num_font": {"size": 8}})
        ch.set_y_axis({"reverse": True, "num_font": {"size": 8}})
        ch.set_legend({"none": True})
        ch.set_chartarea({"border": {"color": BORDURE}})
        ch.set_size({"width": 470, "height": 330})
        ws.insert_chart("J28", ch, {"x_offset": 10, "object_position": 3})
        ws.write("B44", "Les boutons bleu foncé et « DÉMARRER LE COURS » nécessitent l'activation des macros.",
                 self.f(italic=True, font_size=8, font_color=GRIS_TXT))

    # --------------------------------------------------- COURS DU JOUR
    def cours_du_jour(self):
        larg = [21, 34, 12, 9, 10, 30, 11, 11, 11, 11, 11, 11]
        ws = self.feuille("COURS DU JOUR", "shSaisie", ORANGE,
                          "COURS DU JOUR  —  CAHIER DE TEXTE ET APPEL EN DIRECT",
                          "1) Cliquez sur « DÉMARRER LE COURS » : la matière est retrouvée dans l'emploi du temps. "
                          "2) Indiquez la présence du professeur et le contenu. 3) Faites l'appel (double-clic = P → A → R → E). "
                          "4) ENREGISTRER, puis CLÔTURER en fin de cours.", larg, 13)
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        info = self.f(font_color=GRIS_TXT, italic=True, indent=1)
        self.section(ws, 4, 1, 5, "1 · LA SÉANCE")
        self.section(ws, 4, 7, 12, "2 · CAHIER DE TEXTE")
        lignes = [
            (5, "N° séance (recharger)", None), (6, "Matière (code)", None), (7, "Date", "dd/mm/yyyy"),
            (8, "Heure de début", "hh:mm"), (9, "Heure de fin", "hh:mm"), (10, "Enseignant prévu", None),
            (11, "Présence du professeur", None), (12, "Remplaçant / motif", None), (13, "Type de séance", None),
        ]
        for r, t, nf in lignes:
            ws.write(r, 1, t, lab)
            ws.set_row(r, 21)
            if r in (10,):
                continue
            fmt = self.saisie(num_format=nf, align="center", bold=True) if nf else self.saisie(bold=True, indent=1)
            if r == 12:
                ws.merge_range(r, 2, r, 5, "", fmt)
            else:
                ws.write_blank(r, 2, None, fmt)
        ws.merge_range("D6:F6", "", info)
        ws.write_formula("D6", '=IF(C6="","○ Nouvelle séance (non enregistrée)","✔ Séance enregistrée")', info)
        ws.conditional_format("D6", {"type": "formula", "criteria": '=$C$6<>""',
                                     "format": self.f(font_color="#006100", bold=True)})
        ws.merge_range("D7:F7", "", info)
        ws.write_formula("D7", '=IFERROR(XLOOKUP(LEFT(C7,FIND(" ",C7&" ")-1),tblMatieres[Code],tblMatieres[Intitulé]),"")',
                         self.f(bold=True, font_color=BLEU, indent=1))
        ws.merge_range("D8:F8", "", info)
        ws.write_formula("D8", '=IF(C8="","",CHOOSE(WEEKDAY(C8,2),"Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"))', info)
        ws.merge_range("D9:F9", "Format hh:mm (ex. 08:00)", info)
        ws.merge_range("D10:F10", "", info)
        ws.write_formula("D10", '=IF(OR(C9="",C10=""),"","Durée : "&TEXT(MOD(C10-C9,1),"[h]:mm"))', info)
        ws.merge_range("C11:F11", "", self.f(bold=True, indent=1, bg_color=PALE, border=1, border_color=BORDURE))
        ws.write_formula("C11", '=IFERROR(XLOOKUP(LEFT(C7,FIND(" ",C7&" ")-1),tblMatieres[Code],tblMatieres[Enseignant(s)]),"")',
                         self.f(bold=True, indent=1, bg_color=PALE, border=1, border_color=BORDURE))
        ws.merge_range("D12:F12", "", info)
        ws.merge_range("D14:F14", "", info)
        self.dv_liste(ws, "C6", "=LstSeances", ("Recharger une séance", "Choisissez un n° pour la modifier."))
        self.dv_liste(ws, "C7", "=LstMatieres", ("Matière", "Choisissez la matière (remplie par DÉMARRER LE COURS)."))
        ws.data_validation("C8", {"validate": "date", "criteria": ">", "value": dt.date(2020, 1, 1)})
        ws.data_validation("C9:C10", {"validate": "time", "criteria": "between", "minimum": dt.time(0, 0),
                                      "maximum": dt.time(23, 59), "error_message": "Saisissez une heure hh:mm."})
        self.dv_liste(ws, "C12", "=LstProf")
        self.dv_liste(ws, "C14", "=LstTypes")
        self.dv_liste(ws, "C13", "=LstEnseignants", ("Remplaçant", "Choisissez l'enseignant remplaçant ou écrivez un motif."),
                      libre=True)
        ws.conditional_format("C12", {"type": "cell", "criteria": "==", "value": '"Absent"',
                                      "format": self.f(bg_color=STATUTS["A"][0], font_color=STATUTS["A"][1])})
        ws.conditional_format("C12", {"type": "cell", "criteria": "==", "value": '"Présent"',
                                      "format": self.f(bg_color=STATUTS["P"][0], font_color=STATUTS["P"][1])})
        ws.conditional_format("C12", {"type": "cell", "criteria": "==", "value": '"Retard"',
                                      "format": self.f(bg_color=STATUTS["R"][0], font_color=STATUTS["R"][1])})
        # Cahier de texte
        ws.merge_range("H6:M6", "Contenu / chapitres traités", lab)
        ws.merge_range("H7:M10", "", self.saisie(text_wrap=True, valign="top"))
        ws.merge_range("H11:M11", "Travaux demandés / devoirs / prochaine séance", lab)
        ws.merge_range("H12:M13", "", self.saisie(text_wrap=True, valign="top"))
        ws.write("H14", "Observations", lab)
        ws.merge_range("I14:M14", "", self.saisie(indent=1))
        # Boutons
        ws.set_row(15, 42)
        self.rangee(ws, 15, [
            ("▶  DÉMARRER LE COURS", dict(macro="DemarrerCours", couleur=ORANGE, larg=180)),
            ("✔  ENREGISTRER", dict(macro="EnregistrerSeance", couleur=VERT, larg=150)),
            ("■  CLÔTURER (fin = maintenant)", dict(macro="CloturerSeance", couleur=MARINE, larg=210)),
            ("✓  TOUS PRÉSENTS", dict(macro="TousPresents", couleur=BLEU, larg=150)),
            ("＋  NOUVELLE SÉANCE", dict(macro="NouvelleSeance", couleur="#7F7F7F", larg=160)),
            ("📖  Cahier de texte", dict(cible="CAHIER DE TEXTE", couleur=BLEU, larg=150)),
        ], haut=34, dy=4)
        # Appel
        self.section(ws, 17, 1, 12, "3 · APPEL DES ÉTUDIANTS")
        ws.write("B19", "Afficher le groupe", lab)
        ws.write_blank("C19", None, self.saisie(bold=True, indent=1))
        self.dv_liste(ws, "C19", "=LstGroupes", ("Filtrer l'appel", "Laissez vide pour afficher tous les étudiants."))
        ws.merge_range("D19:F19", "(vide = tous les groupes)", info)
        ws.merge_range("B20:F20", "Double-cliquez sur une case « Statut » pour la faire défiler : P → A → R → E", info)
        compteurs = [("Présents", 'COUNTIF($E$22:$E$221,"P")', STATUTS["P"]),
                     ("Absents", 'COUNTIF($E$22:$E$221,"A")', STATUTS["A"]),
                     ("Retards", 'COUNTIF($E$22:$E$221,"R")', STATUTS["R"]),
                     ("Excusés", 'COUNTIF($E$22:$E$221,"E")', STATUTS["E"]),
                     ("Non saisis", "COUNTA($C$22:$C$221)-COUNTA($E$22:$E$221)", ("#FFFFFF", MARINE)),
                     ("Taux", 'IFERROR((H20+J20)/(H20+I20+J20+K20),"")', (BLEU_CLAIR, MARINE))]
        for i, (t, fml, (fond, txt)) in enumerate(compteurs):
            ws.write(18, 7 + i, t, self.f(bold=True, font_size=8, align="center", bg_color=fond, font_color=txt,
                                          border=1, border_color=BORDURE))
            ws.write_formula(19, 7 + i, "=" + fml, self.f(bold=True, font_size=14, align="center", bg_color=fond,
                                                         font_color=txt, border=1, border_color=BORDURE,
                                                         num_format="0%" if t == "Taux" else "0"))
        ws.set_row(19, 26)
        th = self.f(bold=True, font_color=BLANC, bg_color=BLEU, align="center", border=1, border_color=BLANC)
        for c, t in enumerate(["N°", "Nom et prénoms", "Groupe", "Statut", "Arrivée", "Observation"]):
            ws.write(20, 1 + c, t, th)
        ws.set_row(20, 22)
        for r in range(21, 221):
            fond = PALE if r % 2 == 0 else BLANC
            b = {"border": 1, "border_color": "#D9D9D9", "bg_color": fond}
            ws.write_blank(r, 1, None, self.f(align="center", font_color=GRIS_TXT, **b))
            ws.write_blank(r, 2, None, self.f(indent=1, **b))
            ws.write_blank(r, 3, None, self.f(align="center", **b))
            ws.write_blank(r, 4, None, self.f(locked=False, align="center", bold=True, **b))
            ws.write_blank(r, 5, None, self.f(locked=False, align="center", num_format="hh:mm", **b))
            ws.write_blank(r, 6, None, self.f(locked=False, indent=1, **b))
        self.dv_liste(ws, "E22:E221", "=LstStatuts")
        self.cf_statuts(ws, "E22:E221")
        ws.conditional_format("C22:C221", {"type": "formula", "criteria": '=$E22="A"',
                                           "format": self.f(font_color=ROUGE, bold=True)})
        # Légende
        for i, (s, lib) in enumerate(SIGNIF.items()):
            fond, txt = STATUTS[s]
            ws.write(21 + i, 8, s, self.f(bold=True, align="center", bg_color=fond, font_color=txt))
            ws.merge_range(21 + i, 9, 21 + i, 10, lib, self.f(font_color=GRIS_TXT, indent=1))
        ws.freeze_panes(21, 0)
        ws.set_selection("C7")

    # --------------------------------------------------- CAHIER DE TEXTE
    def cahier(self):
        cols = [
            ("ID", 8, None, False, None), ("Date", 11, "dd/mm/yyyy", True, None),
            ("Jour", 10, None, False, '=IF([@Date]="","",CHOOSE(WEEKDAY([@Date],2),"Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"))'),
            ("Début", 8, "hh:mm", True, None), ("Fin", 8, "hh:mm", True, None),
            ("Durée (h)", 8, "0.0", False, '=IF(OR([@Début]="",[@Fin]=""),"",ROUND(MOD([@Fin]-[@Début],1)*24,2))'),
            ("Code matière", 11, None, True, None),
            ("Matière", 30, None, False, '=IF([@[Code matière]]="","",XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Intitulé],"?"))'),
            ("Enseignant prévu", 20, None, False, '=IF([@[Code matière]]="","",XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Enseignant(s)],""))'),
            ("Présence prof", 12, None, True, None), ("Remplaçant ou motif", 18, None, True, None),
            ("Type de séance", 11, None, True, None), ("Contenu / chapitres traités", 50, None, True, None),
            ("Travaux demandés", 30, None, True, None),
            ("Présents", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblPresences[ID séance],[@ID],tblPresences[Statut],"P"))'),
            ("Absents", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblPresences[ID séance],[@ID],tblPresences[Statut],"A"))'),
            ("Retards", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblPresences[ID séance],[@ID],tblPresences[Statut],"R"))'),
            ("Excusés", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblPresences[ID séance],[@ID],tblPresences[Statut],"E"))'),
            ("Taux de présence", 10, "0%", False, '=IF([@ID]="","",IFERROR(([@Présents]+[@Retards])/([@Présents]+[@Absents]+[@Retards]+[@Excusés]),""))'),
            ("Observations", 25, None, True, None),
        ]
        ws = self.feuille("CAHIER DE TEXTE", "shCahier", ORANGE, "CAHIER DE TEXTE  —  TOUTES LES SÉANCES",
                          "Une ligne par séance, créée automatiquement depuis COURS DU JOUR. Les colonnes claires sont modifiables "
                          "(corrections) ; les colonnes calculées sont protégées.", [c[1] for c in cols], len(cols))
        self.rangee(ws, 3, [
            ("✎  Cours du jour", dict(cible="COURS DU JOUR", couleur=ORANGE, larg=140)),
            ("🖨  Export PDF", dict(cible="EXPORT PDF", couleur=VERT, larg=130)),
            ("✖  Supprimer la sélection", dict(macro="SupprimerLignes", couleur=ROUGE, larg=180)),
        ], haut=26, dy=2)
        ws.set_row(3, 26)
        donnees = self.demo_seances() if self.demo else None
        colonnes = []
        for nom, w, nf, saisie, fml in cols:
            kw = {"num_format": nf} if nf else {}
            if nom in ("Contenu / chapitres traités", "Travaux demandés", "Observations"):
                kw["text_wrap"] = True
            if nom in ("ID", "Date", "Jour", "Début", "Fin", "Durée (h)", "Code matière", "Présence prof",
                       "Type de séance", "Présents", "Absents", "Retards", "Excusés", "Taux de présence"):
                kw["align"] = "center"
            fmt = self.f(locked=not saisie, **kw)
            c = {"header": nom, "format": fmt}
            if fml:
                c["formula"] = fml
            colonnes.append(c)
        self.table(ws, 5, 1, "tblSeances", colonnes, donnees)
        n = NB_LIGNES_TABLE
        self.dv_liste(ws, f"K7:K{n}", "=LstProf")
        self.dv_liste(ws, f"M7:M{n}", "=LstTypes")
        self.dv_liste(ws, f"L7:L{n}", "=LstEnseignants", libre=True)
        self.dv_liste(ws, f"H7:H{n}", "=LstMatieres")
        for v, s in (("Absent", "A"), ("Retard", "R"), ("Présent", "P"), ("Remplacé", "E")):
            ws.conditional_format(f"K7:K{n}", {"type": "cell", "criteria": "==", "value": f'"{v}"',
                                               "format": self.f(bg_color=STATUTS[s][0], font_color=STATUTS[s][1])})
        ws.conditional_format(f"T7:T{n}", {"type": "data_bar", "bar_color": "#9BC2E6", "bar_solid": True,
                                           "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 1})
        ws.freeze_panes(6, 3)

    # ------------------------------------------------------- PRÉSENCES
    def presences(self):
        cols = [("ID séance", 9, None, False), ("Étudiant", 32, None, False), ("Groupe", 10, None, False),
                ("Code matière", 12, None, False), ("Matière", 30, None, False), ("Date", 11, "dd/mm/yyyy", False),
                ("Heure", 8, "hh:mm", False), ("Statut", 8, None, True), ("Arrivée", 9, "hh:mm", True),
                ("Observation", 28, None, True), ("Saisie le", 16, "dd/mm/yyyy hh:mm", False)]
        ws = self.feuille("PRÉSENCES", "shPresences", ORANGE, "REGISTRE DES PRÉSENCES",
                          "Rempli automatiquement par COURS DU JOUR (une ligne par étudiant et par séance). Pour corriger, "
                          "rechargez la séance dans COURS DU JOUR, ou modifiez directement Statut / Arrivée / Observation.",
                          [c[1] for c in cols], len(cols))
        colonnes = []
        for nom, w, nf, saisie in cols:
            kw = {"num_format": nf} if nf else {}
            if nom not in ("Étudiant", "Matière", "Observation"):
                kw["align"] = "center"
            colonnes.append({"header": nom, "format": self.f(locked=not saisie, **kw)})
        self.table(ws, 5, 1, "tblPresences", colonnes, self.demo_presences() if self.demo else None,
                   style="Table Style Medium 9")
        n = 60000
        self.dv_liste(ws, f"I7:I{n}", "=LstStatuts")
        self.cf_statuts(ws, f"I7:I{n}")
        ws.freeze_panes(6, 3)

    # --------------------------------------------------------- EXPOSÉS
    def exposes(self):
        larg = [12, 30, 10, 30, 12, 12, 11, 9, 25, 22]
        ws = self.feuille("EXPOSÉS", "shExposes", ORANGE, "EXPOSÉS  —  GROUPES PASSÉS ET RESTANTS",
                          "Choisissez une matière puis « Une ligne par groupe ». Après chaque passage, sélectionnez la ligne "
                          "et cliquez sur « Marquer présenté ». Le tableau de synthèse se met à jour seul.", larg, 10)
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        ws.write("B4", "Matière", lab)
        ws.write_blank("C4", None, self.saisie(bold=True, indent=1))
        self.dv_liste(ws, "C4", "=LstMatieres")
        ws.set_row(3, 30)
        self.rangee(ws, 3, [
            ("⇩  Une ligne par groupe", dict(macro="GenererExposes", couleur=ORANGE, larg=170)),
            ("✔  Marquer présenté", dict(macro="MarquerPresente", couleur=VERT, larg=150)),
            ("Restants", dict(macro="AfficherExposesRestants", couleur=BLEU, larg=85)),
            ("Tout afficher", dict(macro="ToutAfficherExposes", couleur=BLEU, larg=100)),
            ("＋ Ligne", dict(macro="AjouterLigne", couleur="#7F7F7F", larg=80)),
            ("✖ Supprimer", dict(macro="SupprimerLignes", couleur=ROUGE, larg=100)),
        ], x0=330)
        # Synthèse
        self.section(ws, 5, 1, 10, "SYNTHÈSE PAR MATIÈRE")
        th = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", text_wrap=True, bottom=1)
        for c, t in [(1, "Code"), (2, "Matière"), (3, "Groupes"), (4, "Ont présenté"), (5, "Restants")]:
            ws.write(6, c, t, th)
        ws.merge_range("G7:H7", "Groupes ayant présenté", th)
        ws.merge_range("I7:K7", "Groupes restants", th)
        ws.set_row(6, 28)
        for k in range(15):
            r = 7 + k
            x = r + 1
            fond = PALE if k % 2 else BLANC
            cel = self.f(bg_color=fond, align="center")
            celg = self.f(bg_color=fond, indent=1, text_wrap=True)
            ws.write_formula(r, 1, f'=IFERROR(INDEX(tblMatieres[Code],{k + 1}),"")', self.f(bg_color=fond, bold=True,
                                                                                         align="center"))
            ws.write_formula(r, 2, f'=IF($B{x}="","",INDEX(tblMatieres[Intitulé],{k + 1}))', celg)
            ws.write_blank(r, 3, None, cel)
            ws.write_formula(r, 4, f'=IF($B{x}="","",IF(COUNTIF(tblExposes[Code matière],$B{x})=0,"—",COUNTIFS(tblExposes[Code matière],$B{x},tblExposes[Statut],"Présenté")))', cel)
            ws.write_blank(r, 5, None, cel)
            ws.merge_range(r, 6, r, 7, "", celg)
            ws.merge_range(r, 8, r, 10, "", celg)
        ws.conditional_format("F8:F22", {"type": "cell", "criteria": ">", "value": 0,
                                         "format": self.f(bold=True, font_color="#C55A11")})
        ws.conditional_format("I8:I22", {"type": "formula", "criteria": '=LEFT($I8,1)="✔"',
                                         "format": self.f(bold=True, font_color="#006100")})
        # Tableau des exposés
        self.section(ws, 23, 1, 10, "LISTE DES EXPOSÉS")
        cols = [("Code matière", None, True, None), ("Matière", None, False,
                 '=IF([@[Code matière]]="","",XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Abrégé],"?"))'),
                ("Groupe", None, True, None), ("Thème", None, True, None), ("Date prévue", "dd/mm/yyyy", True, None),
                ("Date de passage", "dd/mm/yyyy", True, None), ("Statut", None, True, None),
                ("Note sur 20", "0.0", True, None), ("Observation", None, True, None), ("Membres", None, False, None)]
        colonnes = []
        for nom, nf, saisie, fml in cols:
            kw = {"num_format": nf} if nf else {}
            if nom in ("Code matière", "Groupe", "Date prévue", "Date de passage", "Statut", "Note sur 20"):
                kw["align"] = "center"
            if nom == "Membres":
                kw.update({"font_size": 7, "font_color": "#A6A6A6"})
            if nom in ("Thème", "Observation"):
                kw["text_wrap"] = True
            c = {"header": nom, "format": self.f(locked=not saisie, **kw)}
            if fml:
                c["formula"] = fml
            colonnes.append(c)
        self.table(ws, 24, 1, "tblExposes", colonnes, style="Table Style Medium 3")
        n = NB_LIGNES_TABLE
        self.dv_liste(ws, f"B26:B{n}", "=LstMatieres")
        self.dv_liste(ws, f"D26:D{n}", "=LstGroupes")
        self.dv_liste(ws, f"H26:H{n}", "=LstStatutExpose")
        for v, s in (("Présenté", "P"), ("Reporté", "R"), ("Annulé", "E"), ("À venir", None)):
            fmt = self.f(bg_color=STATUTS[s][0], font_color=STATUTS[s][1], bold=True) if s else \
                self.f(bg_color=BLEU_CLAIR, font_color=MARINE)
            ws.conditional_format(f"H26:H{n}", {"type": "cell", "criteria": "==", "value": f'"{v}"', "format": fmt})

    # -------------------------------------------------------- RAPPORTS
    def rapports(self):
        cols = [
            ("ID", 7, None, False, None), ("Code matière", 12, None, True, None),
            ("Matière", 18, None, False, '=IF([@[Code matière]]="","",XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Abrégé],"?"))'),
            ("Intitulé", 36, None, True, None), ("Type", 11, None, True, None),
            ("Date consigne", 11, "dd/mm/yyyy", True, None), ("Date limite", 11, "dd/mm/yyyy", True, None),
            ("Heure limite", 9, "hh:mm", True, None),
            ("Attendus", 9, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID]))'),
            ("Remis", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"À temps")+COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"En retard")+COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"Remis"))'),
            ("À temps", 8, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"À temps"))'),
            ("En retard", 9, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"En retard"))'),
            ("Non remis", 9, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"Non remis"))'),
            ("En attente", 9, "0", False, '=IF([@ID]="","",COUNTIFS(tblRemises[ID rapport],[@ID],tblRemises[Statut],"En attente"))'),
            ("Taux de remise", 12, "0%", False, '=IF([@ID]="","",IFERROR([@Remis]/[@Attendus],""))'),
        ]
        ws = self.feuille("RAPPORTS", "shRapports", ORANGE, "RAPPORTS  —  CONSIGNES ET BILAN DES REMISES",
                          "« Nouveau rapport » crée une consigne. Sélectionnez ensuite sa ligne puis « Générer les remises "
                          "attendues » : une ligne par étudiant (Individuel) ou par groupe (Groupe) est créée dans REMISES.",
                          [c[1] for c in cols], len(cols))
        ws.set_row(3, 30)
        self.rangee(ws, 3, [
            ("＋  Nouveau rapport", dict(macro="NouveauRapport", couleur=ORANGE, larg=160)),
            ("⇩  Générer les remises attendues", dict(macro="GenererRemises", couleur=VERT, larg=240)),
            ("📥  Feuille REMISES", dict(cible="REMISES", couleur=BLEU, larg=150)),
            ("✖  Supprimer", dict(macro="SupprimerLignes", couleur=ROUGE, larg=110)),
        ])
        colonnes = []
        for nom, w, nf, saisie, fml in cols:
            kw = {"num_format": nf} if nf else {}
            if nom not in ("Matière", "Intitulé"):
                kw["align"] = "center"
            if nom == "Intitulé":
                kw["text_wrap"] = True
            c = {"header": nom, "format": self.f(locked=not saisie, **kw)}
            if fml:
                c["formula"] = fml
            colonnes.append(c)
        self.table(ws, 5, 1, "tblRapports", colonnes, style="Table Style Medium 4")
        n = NB_LIGNES_TABLE
        self.dv_liste(ws, f"C7:C{n}", "=LstMatieres")
        self.dv_liste(ws, f"F7:F{n}", "=LstTypeRapport")
        ws.conditional_format(f"N7:N{n}", {"type": "cell", "criteria": ">", "value": 0,
                                           "format": self.f(bg_color=STATUTS["A"][0], font_color=STATUTS["A"][1],
                                                            bold=True)})
        ws.conditional_format(f"P7:P{n}", {"type": "data_bar", "bar_color": "#A9D08E", "bar_solid": True,
                                           "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 1})
        ws.freeze_panes(6, 2)

    def remises(self):
        cols = [
            ("ID rapport", 10, None, True, None),
            ("Code matière", 12, None, False, '=IF([@[ID rapport]]="","",XLOOKUP([@[ID rapport]],tblRapports[ID],tblRapports[Code matière],""))'),
            ("Intitulé", 32, None, False, '=IF([@[ID rapport]]="","",XLOOKUP([@[ID rapport]],tblRapports[ID],tblRapports[Intitulé],""))'),
            ("Type", 10, None, False, '=IF([@[ID rapport]]="","",XLOOKUP([@[ID rapport]],tblRapports[ID],tblRapports[Type],""))'),
            ("Remis par", 30, None, True, None),
            ("Échéance", 15, "dd/mm/yyyy hh:mm", False, '=IF([@[ID rapport]]="","",XLOOKUP([@[ID rapport]],tblRapports[ID],tblRapports[Date limite],0)+XLOOKUP([@[ID rapport]],tblRapports[ID],tblRapports[Heure limite],0))'),
            ("Date remise", 11, "dd/mm/yyyy", True, None), ("Heure remise", 9, "hh:mm", True, None),
            ("Statut", 11, None, False, '=IF([@[ID rapport]]="","",IF([@[Date remise]]="",IF(AND([@Échéance]>1,NOW()>[@Échéance]),"Non remis","En attente"),IF([@Échéance]<1,"Remis",IF([@[Date remise]]+[@[Heure remise]]<=[@Échéance]+1/1440,"À temps","En retard"))))'),
            ("Observation", 25, None, True, None), ("Membres", 22, None, False, None),
        ]
        ws = self.feuille("REMISES", "shRemises", ORANGE, "REMISES DES RAPPORTS  —  QUI A REMIS, QUAND",
                          "Sélectionnez la ou les lignes concernées puis « Marquer remis maintenant » : la date et l'heure du "
                          "dépôt sont enregistrées et le statut (À temps / En retard / Non remis) se calcule seul.",
                          [c[1] for c in cols], len(cols))
        ws.set_row(3, 30)
        self.rangee(ws, 3, [
            ("✔  Marquer remis maintenant", dict(macro="MarquerRemis", couleur=VERT, larg=210)),
            ("Non remis", dict(macro="AfficherNonRemis", couleur=ROUGE, larg=95)),
            ("En retard", dict(macro="AfficherEnRetard", couleur="#C55A11", larg=95)),
            ("Tout afficher", dict(macro="ToutAfficherRemises", couleur=BLEU, larg=105)),
            ("＋ Ligne", dict(macro="AjouterLigne", couleur="#7F7F7F", larg=80)),
            ("✖ Supprimer", dict(macro="SupprimerLignes", couleur=ROUGE, larg=105)),
            ("📄 Rapports", dict(cible="RAPPORTS", couleur=BLEU, larg=105)),
        ])
        colonnes = []
        for nom, w, nf, saisie, fml in cols:
            kw = {"num_format": nf} if nf else {}
            if nom not in ("Intitulé", "Remis par", "Observation", "Membres"):
                kw["align"] = "center"
            if nom == "Membres":
                kw.update({"font_size": 7, "font_color": "#A6A6A6"})
            c = {"header": nom, "format": self.f(locked=not saisie, **kw)}
            if fml:
                c["formula"] = fml
            colonnes.append(c)
        self.table(ws, 5, 1, "tblRemises", colonnes, style="Table Style Medium 4")
        n = NB_LIGNES_TABLE
        self.dv_liste(ws, f"B7:B{n}", "=LstRapports")
        self.dv_liste(ws, f"F7:F{n}", "=LstEtudiants", ("Remis par", "Étudiant (rapport individuel) ou nom du groupe."),
                      libre=True)
        for v, s in (("À temps", "P"), ("Remis", "P"), ("En retard", "R"), ("Non remis", "A"), ("En attente", "E")):
            ws.conditional_format(f"J7:J{n}", {"type": "cell", "criteria": "==", "value": f'"{v}"',
                                               "format": self.f(bg_color=STATUTS[s][0], font_color=STATUTS[s][1],
                                                                bold=True)})
        ws.freeze_panes(6, 2)

    # --------------------------------------------------- FICHE ÉTUDIANT
    def fiche(self):
        larg = [16, 30, 10, 8, 7, 7, 7, 7, 9, 12, 3,
                12, 30, 10, 8, 8, 3,
                12, 28, 10, 22, 15, 11, 8, 11, 3,
                12, 18, 10, 28, 11, 11, 10, 7]
        ws = self.feuille("FICHE ÉTUDIANT", "shFiche", VERT, "FICHE ÉTUDIANT  —  PRÉSENCES, ABSENCES, RAPPORTS ET EXPOSÉS",
                          "Tapez un nom (ou une partie) dans « Rechercher » puis Entrée, ou choisissez l'étudiant dans la "
                          "liste. Tout se met à jour automatiquement.", larg, 10)
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        info = self.f(font_color=GRIS_TXT, italic=True, indent=1)
        nom = "$C$7"
        ws.write("B5", "🔎 Rechercher", lab)
        ws.write_blank("C5", None, self.saisie(indent=1))
        ws.merge_range("D5:J5", "← tapez « akpa », « brice »… puis Entrée (macros activées)", info)
        ws.write("B7", "Étudiant", lab)
        ws.write_blank("C7", None, self.saisie(bold=True, font_size=12, indent=1, font_color=MARINE))
        self.dv_liste(ws, "C7", "=LstEtudiants", ("Étudiant", "Choisissez l'étudiant dans la liste."))
        ws.set_row(6, 24)
        self.rangee(ws, 6, [
            ("🖨  Fiche en PDF", dict(macro="ExporterFichePDF", couleur=VERT, larg=130)),
            ("↻  Actualiser", dict(macro="Actualiser", couleur=BLEU, larg=100)),
        ], haut=24, dy=4, x0=330)
        vals = self.f(indent=1, bg_color=PALE, border=1, border_color=BORDURE)
        ws.write("B9", "Groupe (défaut)", lab)
        ws.write_formula("C9", f'=IF({nom}="","",XLOOKUP({nom},tblEtudiants[Nom et prénoms],tblEtudiants[Groupe],""))', vals)
        ws.write("B10", "Rôle", lab)
        ws.write_formula("C10", f'=IF({nom}="","",XLOOKUP({nom},tblEtudiants[Nom et prénoms],tblEtudiants[Rôle],"")&"")', vals)
        ws.write("D9", "Email", lab)
        ws.merge_range("E9:J9", "", vals)
        ws.write_formula("E9", f'=IF({nom}="","",XLOOKUP({nom},tblEtudiants[Nom et prénoms],tblEtudiants[Email],"")&"")', vals)
        ws.write("D10", "Contact", lab)
        ws.merge_range("E10:J10", "", vals)
        ws.write_formula("E10", f'=IF({nom}="","",XLOOKUP({nom},tblEtudiants[Nom et prénoms],tblEtudiants[Contact],""))',
                         self.f(indent=1, bg_color=PALE, border=1, border_color=BORDURE, num_format="00 00 00 00 00",
                                align="left"))
        # Bilan
        cnt = lambda s: f'COUNTIFS(tblPresences[Étudiant],{nom},tblPresences[Statut],"{s}")'
        cartes = [
            ("B", "B", "SÉANCES", f'=IF({nom}="","",COUNTIF(tblPresences[Étudiant],{nom}))', BLEU, "0"),
            ("C", "C", "TAUX DE PRÉSENCE", f'=IF({nom}="","",IFERROR(({cnt("P")}+{cnt("R")})/COUNTIF(tblPresences[Étudiant],{nom}),"—"))', MARINE, "0%"),
            ("D", "E", "PRÉSENT", f'=IF({nom}="","",{cnt("P")})', "#2E7D32", "0"),
            ("F", "G", "ABSENT", f'=IF({nom}="","",{cnt("A")})', ROUGE, "0"),
            ("H", "I", "RETARD", f'=IF({nom}="","",{cnt("R")})', "#C55A11", "0"),
            ("J", "J", "EXCUSÉ", f'=IF({nom}="","",{cnt("E")})', "#7F7F7F", "0"),
        ]
        ws.set_row(12, 34)
        for c1, c2, lib, fml, coul, nf in cartes:
            t = self.f(bold=True, font_size=8, font_color=BLANC, bg_color=coul, align="center")
            v = self.f(bold=True, font_size=18, font_color=coul, bg_color=PALE, align="center", num_format=nf)
            if c1 == c2:
                ws.write(f"{c1}12", lib, t)
                ws.write_formula(f"{c1}13", fml, v)
            else:
                ws.merge_range(f"{c1}12:{c2}12", lib, t)
                ws.merge_range(f"{c1}13:{c2}13", "", v)
                ws.write_formula(f"{c1}13", fml, v)
        # Par matière
        self.section(ws, 14, 1, 10, "PAR MATIÈRE  (rouge : plus de « seuil » absences ; orange : seuil atteint)")
        th = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1)
        for c, t in enumerate(["Code", "Matière", "Groupe", "Séances", "P", "A", "R", "E", "Taux", "Alerte"]):
            ws.write(15, 1 + c, t, th)
        for k in range(15):
            r = 16 + k
            x = r + 1
            fond = PALE if k % 2 else BLANC
            cel = self.f(bg_color=fond, align="center")
            code = f"$B{x}"
            vide = f'OR({nom}="",{code}="")'
            ws.write_formula(r, 1, f'=IFERROR(INDEX(tblMatieres[Code],{k + 1}),"")', self.f(bg_color=fond, bold=True,
                                                                                         align="center"))
            ws.write_formula(r, 2, f'=IF({code}="","",INDEX(tblMatieres[Intitulé],{k + 1}))',
                             self.f(bg_color=fond, indent=1))
            ws.write_dynamic_array_formula(
                r, 3, r, 3,
                f'=IF({vide},"",IF(COUNTIF(tblGroupes[Code matière],{code})>0,XLOOKUP(1,(tblGroupes[Code matière]={code})*(tblGroupes[Étudiant]={nom}),tblGroupes[Groupe],"—"),XLOOKUP({nom},tblEtudiants[Nom et prénoms],tblEtudiants[Groupe],"")))',
                cel)
            base = f"tblPresences[Étudiant],{nom},tblPresences[Code matière],{code}"
            ws.write_formula(r, 4, f'=IF({vide},"",COUNTIFS({base}))', cel)
            for j, s in enumerate("PARE"):
                ws.write_formula(r, 5 + j, f'=IF({vide},"",COUNTIFS({base},tblPresences[Statut],"{s}"))', cel)
            ws.write_formula(r, 9, f'=IF({vide},"",IF(E{x}=0,"",(F{x}+H{x})/E{x}))',
                             self.f(bg_color=fond, align="center", num_format="0%"))
            ws.write_formula(r, 10, f'=IF({vide},"",IF(G{x}>pSeuil,"⚠ ALERTE",IF(G{x}=pSeuil,"Attention","")))', cel)
        ws.conditional_format("G17:G31", {"type": "formula", "criteria": "=AND(ISNUMBER(G17),G17>pSeuil)",
                                          "format": self.f(bg_color=ROUGE, font_color=BLANC, bold=True)})
        ws.conditional_format("G17:G31", {"type": "formula", "criteria": "=AND(ISNUMBER(G17),G17=pSeuil)",
                                          "format": self.f(bg_color="#F4B183", font_color="#843C0C", bold=True)})
        ws.conditional_format("K17:K31", {"type": "formula", "criteria": '=LEFT($K17,1)="⚠"',
                                          "format": self.f(font_color=ROUGE, bold=True)})
        # Listes dynamiques
        tht = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1, text_wrap=True)
        self.section(ws, 32, 1, 4, "ABSENCES", ROUGE)
        for c, t in enumerate(["Code", "Matière", "Date", "Heure"]):
            ws.write(33, 1 + c, t, tht)
        self.section(ws, 4, 12, 16, "HISTORIQUE COMPLET DES SÉANCES", BLEU)
        for c, t in enumerate(["Code", "Matière", "Date", "Heure", "Statut"]):
            ws.write(5, 12 + c, t, tht)
        self.section(ws, 4, 18, 25, "RAPPORTS (individuels et de ses groupes)", ORANGE)
        for c, t in enumerate(["Code", "Intitulé", "Type", "Remis par", "Échéance", "Date remise", "Heure", "Statut"]):
            ws.write(5, 18 + c, t, tht)
        self.section(ws, 4, 27, 34, "EXPOSÉS DE SES GROUPES", OR)
        for c, t in enumerate(["Code", "Matière", "Groupe", "Thème", "Date prévue", "Passage", "Statut", "Note"]):
            ws.write(5, 27 + c, t, tht)
        # Formats des colonnes de listes
        for col, nf in (("D", "dd/mm/yyyy"), ("E", "hh:mm"), ("O", "dd/mm/yyyy"), ("P", "hh:mm"),
                        ("W", "dd/mm/yyyy hh:mm"), ("X", "dd/mm/yyyy"), ("Y", "hh:mm"),
                        ("AF", "dd/mm/yyyy"), ("AG", "dd/mm/yyyy")):
            debut = 35 if col in ("D", "E") else 7
            for r in range(debut - 1, debut + 299):
                ws.write_blank(r, ord(col[-1]) - 65 + (26 if len(col) == 2 else 0), None,
                               self.f(num_format=nf, align="center", font_size=9))
        self.cf_statuts(ws, "Q7:Q400")
        for v, s in (("À temps", "P"), ("Remis", "P"), ("En retard", "R"), ("Non remis", "A")):
            ws.conditional_format("Z7:Z400", {"type": "cell", "criteria": "==", "value": f'"{v}"',
                                              "format": self.f(bg_color=STATUTS[s][0], font_color=STATUTS[s][1])})
        ws.conditional_format("AH7:AH400", {"type": "cell", "criteria": "==", "value": '"Présenté"',
                                            "format": self.f(bg_color=STATUTS["P"][0], font_color=STATUTS["P"][1])})
        if not self.apercu:
            ws.set_landscape()
            ws.set_paper(9)
            ws.fit_to_pages(1, 0)
            ws.print_area("B2:Q60")
        ws.set_selection("C5")

    # ------------------------------------------------ SUIVI ENSEIGNANTS
    def suivi(self):
        larg = [12, 30, 22, 10, 10, 10, 10, 11, 10, 14, 13, 45, 12]
        ws = self.feuille("SUIVI ENSEIGNANTS", "shSuivi", VERT, "SUIVI DES ENSEIGNANTS ET DE L'AVANCEMENT DES COURS",
                          "Calculé automatiquement à partir du cahier de texte. Les heures d'une séance sans professeur ne "
                          "sont pas comptées.", larg, 13)
        self.section(ws, 4, 1, 13, "AVANCEMENT PAR MATIÈRE")
        th = self.f(bold=True, font_color=BLANC, bg_color=BLEU, align="center", text_wrap=True, border=1,
                    border_color=BLANC)
        entetes = ["Code", "Matière", "Enseignant(s)", "Volume prévu (h)", "Séances tenues", "Prof absent",
                   "Prof en retard", "Heures effectuées", "Reste (h)", "Avancement", "Dernière séance",
                   "Dernier contenu traité", "Présence étudiants"]
        for c, t in enumerate(entetes):
            ws.write(5, 1 + c, t, th)
        ws.set_row(5, 32)
        for k in range(15):
            r = 6 + k
            x = r + 1
            fond = PALE if k % 2 else BLANC
            cel = self.f(bg_color=fond, align="center")
            c = f"$B{x}"
            vide = f'{c}=""'
            ws.write_formula(r, 1, f'=IFERROR(INDEX(tblMatieres[Code],{k + 1}),"")', self.f(bg_color=fond, bold=True,
                                                                                         align="center"))
            ws.write_formula(r, 2, f'=IF({vide},"",INDEX(tblMatieres[Abrégé],{k + 1}))', self.f(bg_color=fond, indent=1))
            ws.write_formula(r, 3, f'=IF({vide},"",INDEX(tblMatieres[Enseignant(s)],{k + 1}))',
                             self.f(bg_color=fond, indent=1))
            ws.write_formula(r, 4, f'=IF({vide},"",INDEX(tblMatieres[Volume total (h)],{k + 1}))', cel)
            ws.write_formula(r, 5, f'=IF({vide},"",COUNTIFS(tblSeances[Code matière],{c},tblSeances[Présence prof],"<>Absent"))', cel)
            ws.write_formula(r, 6, f'=IF({vide},"",COUNTIFS(tblSeances[Code matière],{c},tblSeances[Présence prof],"Absent"))', cel)
            ws.write_formula(r, 7, f'=IF({vide},"",COUNTIFS(tblSeances[Code matière],{c},tblSeances[Présence prof],"Retard"))', cel)
            ws.write_formula(r, 8, f'=IF({vide},"",SUMIFS(tblSeances[Durée (h)],tblSeances[Code matière],{c},tblSeances[Présence prof],"<>Absent"))',
                             self.f(bg_color=fond, align="center", num_format="0.0", bold=True))
            ws.write_formula(r, 9, f'=IF({vide},"",MAX(0,E{x}-I{x}))', self.f(bg_color=fond, align="center",
                                                                              num_format="0.0"))
            ws.write_formula(r, 10, f'=IF({vide},"",IF(N(E{x})=0,0,I{x}/E{x}))',
                             self.f(bg_color=fond, align="center", num_format="0%"))
            ws.write_array_formula(r, 11, r, 11, f'{{=IF({vide},"",IF(COUNTIF(tblSeances[Code matière],{c})=0,"—",MAX(IF(tblSeances[Code matière]={c},tblSeances[Date]))))}}',
                             self.f(bg_color=fond, align="center", num_format="dd/mm/yyyy"))
            ws.write_array_formula(r, 12, r, 12, f'=IF({vide},"",IFERROR(LOOKUP(2,1/(tblSeances[Code matière]={c}),tblSeances[Contenu / chapitres traités])&"","—"))',
                             self.f(bg_color=fond, indent=1, font_size=9, text_wrap=True))
            ws.write_formula(r, 13, f'=IF({vide},"",IFERROR((COUNTIFS(tblPresences[Code matière],{c},tblPresences[Statut],"P")+COUNTIFS(tblPresences[Code matière],{c},tblPresences[Statut],"R"))/COUNTIFS(tblPresences[Code matière],{c}),"—"))',
                             self.f(bg_color=fond, align="center", num_format="0%"))
        ws.conditional_format("K7:K21", {"type": "data_bar", "bar_color": "#5B9BD5", "bar_solid": True,
                                         "min_type": "num", "min_value": 0, "max_type": "num", "max_value": 1})
        ws.conditional_format("G7:G21", {"type": "cell", "criteria": ">", "value": 0,
                                         "format": self.f(bg_color=STATUTS["A"][0], font_color=STATUTS["A"][1],
                                                          bold=True)})
        ws.conditional_format("H7:H21", {"type": "cell", "criteria": ">", "value": 0,
                                         "format": self.f(bg_color=STATUTS["R"][0], font_color=STATUTS["R"][1])})
        # Absences et retards
        self.section(ws, 23, 1, 13, "ABSENCES, RETARDS ET REMPLACEMENTS DES ENSEIGNANTS", ROUGE)
        th2 = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1)
        for c, t in enumerate(["Date", "Jour", "Début", "Fin", "Durée (h)", "Code", "Matière", "Enseignant prévu",
                               "Présence", "Remplaçant ou motif"]):
            ws.write(24, 1 + c, t, th2)
        for r in range(25, 325):
            ws.write_blank(r, 1, None, self.f(num_format="dd/mm/yyyy", align="center"))
            ws.write_blank(r, 3, None, self.f(num_format="hh:mm", align="center"))
            ws.write_blank(r, 4, None, self.f(num_format="hh:mm", align="center"))
        self.par_enseignant(ws)
        ws.freeze_panes(6, 2)

    def par_enseignant(self, ws):
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        for col, w in zip(range(14, 27), [3, 12, 10, 8, 8, 8, 12, 24, 20, 11, 18, 10, 45]):
            ws.set_column(col, col, w)
        self.section(ws, 4, 15, 26, "SUIVI PAR ENSEIGNANT", OR)
        ws.write("P6", "Enseignant", lab)
        ws.merge_range("Q6:S6", "", self.saisie(bold=True, indent=1))
        self.dv_liste(ws, "Q6", "=LstEnseignants", ("Enseignant", "Choisissez un enseignant dans la liste."))
        ws.merge_range("T6:W6", "← choisissez un enseignant dans la liste", self.f(italic=True, font_color=GRIS_TXT,
                                                                                indent=1))
        e = '"*"&$Q$6&"*"'
        ok = f'tblSeances[Enseignant prévu],{e}'
        cartes = [
            ("P", "Q", "MATIÈRES", f'=IF($Q$6="","",COUNTIF(tblMatieres[Enseignant(s)],{e}))', BLEU, "0"),
            ("R", "S", "SÉANCES TENUES", f'=IF($Q$6="","",COUNTIFS({ok},tblSeances[Présence prof],"<>Absent"))', VERT, "0"),
            ("T", "T", "ABSENCES", f'=IF($Q$6="","",COUNTIFS({ok},tblSeances[Présence prof],"Absent"))', ROUGE, "0"),
            ("U", "U", "RETARDS", f'=IF($Q$6="","",COUNTIFS({ok},tblSeances[Présence prof],"Retard"))', "#C55A11", "0"),
            ("V", "V", "HEURES FAITES", f'=IF($Q$6="","",SUMIFS(tblSeances[Durée (h)],{ok},tblSeances[Présence prof],"<>Absent"))', MARINE, '0.0" h"'),
            ("W", "X", "AVANCEMENT", f'=IF($Q$6="","",IFERROR(V9/SUMIFS(tblMatieres[Volume total (h)],tblMatieres[Enseignant(s)],{e}),0))', OR, "0%"),
        ]
        ws.set_row(8, 30)
        for c1, c2, lib, fml, coul, nf in cartes:
            t = self.f(bold=True, font_size=8, font_color=BLANC, bg_color=coul, align="center")
            v = self.f(bold=True, font_size=16, font_color=coul, bg_color=PALE, align="center", num_format=nf)
            if c1 == c2:
                ws.write(f"{c1}8", lib, t)
                ws.write_formula(f"{c1}9", fml, v)
            else:
                ws.merge_range(f"{c1}8:{c2}8", lib, t)
                ws.merge_range(f"{c1}9:{c2}9", "", v)
                ws.write_formula(f"{c1}9", fml, v)
        ws.write_formula("Y9", f'=IF($Q$6="","","sur "&SUMIFS(tblMatieres[Volume total (h)],tblMatieres[Enseignant(s)],{e})&" h prévues")',
                         self.f(font_color=GRIS_TXT, italic=True, indent=1))
        th = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1, text_wrap=True)
        for i, t in enumerate(["Date", "Jour", "Début", "Fin", "Durée", "Code", "Matière", "Enseignant prévu",
                               "Présence", "Remplaçant", "Type", "Contenu traité"]):
            ws.write(10, 15 + i, t, th)
        for r in range(11, 311):
            ws.write_blank(r, 15, None, self.f(num_format="dd/mm/yyyy", align="center", font_size=9))
            ws.write_blank(r, 17, None, self.f(num_format="hh:mm", align="center", font_size=9))
            ws.write_blank(r, 18, None, self.f(num_format="hh:mm", align="center", font_size=9))
        for v, st in (("Absent", "A"), ("Retard", "R"), ("Présent", "P"), ("Remplacé", "E")):
            ws.conditional_format("X12:X311", {"type": "cell", "criteria": "==", "value": f'"{v}"',
                                               "format": self.f(bg_color=STATUTS[st][0], font_color=STATUTS[st][1])})

    # --------------------------------------------------- MODE D'EMPLOI
    def guide(self):
        ws = self.feuille("MODE D'EMPLOI", "shGuide", OR, "MODE D'EMPLOI  —  COMMENT UTILISER LE CLASSEUR",
                          "Lisez une fois ce guide. Les cases jaune clair sont celles que vous remplissez ; tout le reste "
                          "se calcule seul. Cliquez sur les boutons de couleur pour aller aux feuilles.", [5, 120, 3, 22], 4)
        blocs = [
            ("A", "AVANT DE COMMENCER (une seule fois)", MARINE, None, [
                "Débloquez le fichier : clic droit sur le fichier › Propriétés › cochez « Débloquer » › OK. Ouvrez-le "
                "puis cliquez sur « Activer le contenu » dans le bandeau jaune. Sans cela les boutons ne marchent pas.",
                "Feuille PARAMÈTRES : écrivez votre nom dans « Délégué de classe » (il apparaît sur les PDF).",
                "Feuille ÉTUDIANTS : vérifiez la liste des étudiants (nom, groupe, email, contact). Le bouton "
                "« Ajouter un étudiant » ajoute une ligne.",
                "Les feuilles sont protégées pour éviter les erreurs (mot de passe : GC5BTP, bouton « Mode "
                "administrateur » sur l'accueil). Les listes déroulantes apparaissent quand vous cliquez sur une case "
                "jaune : petite flèche à droite de la case.",
            ]),
            ("B", "À CHAQUE COURS  (feuille COURS DU JOUR)", ORANGE, "COURS DU JOUR", [
                "Sur l'ACCUEIL, cliquez sur le bouton orange « ▶ DÉMARRER LE COURS ». Le classeur regarde le jour, "
                "l'heure et l'emploi du temps : la matière, la date et l'heure de début se remplissent seules. "
                "S'il n'y a pas de cours prévu à cette heure, choisissez la matière dans la liste.",
                "« Présence du professeur » : choisissez Présent, Absent, Retard ou Remplacé dans la liste. "
                "En cas de remplacement, choisissez le remplaçant dans la liste « Remplaçant / motif ».",
                "Écrivez le contenu du cours (chapitres traités) et les travaux demandés dans les grandes cases jaunes.",
                "L'appel : cliquez sur « ✓ TOUS PRÉSENTS », puis double-cliquez sur la case « Statut » d'un étudiant "
                "pour la changer : P → A (absent) → R (retard, l'heure d'arrivée est notée) → E (excusé). "
                "La case « Afficher le groupe » permet de ne voir qu'un groupe.",
                "Cliquez sur « ✔ ENREGISTRER ». La séance va dans le CAHIER DE TEXTE et l'appel dans PRÉSENCES. "
                "Vous pouvez enregistrer plusieurs fois.",
                "En fin de cours, cliquez sur « ■ CLÔTURER » : l'heure de fin est notée et la durée calculée.",
                "Pour corriger une séance passée : choisissez son numéro dans « N° séance (recharger) », modifiez, "
                "puis ENREGISTRER.",
            ]),
            ("C", "EXPOSÉS", "#C55A11", "EXPOSÉS", [
                "Feuille EXPOSÉS : choisissez la matière dans la case jaune « Matière ».",
                "Cliquez sur « Une ligne par groupe » : une ligne est créée pour chaque groupe. Complétez le thème "
                "et la date prévue.",
                "Le jour du passage, cliquez sur la ligne du groupe puis sur « ✔ Marquer présenté ». Le tableau du "
                "haut indique pour chaque matière les groupes passés et ceux qui restent.",
            ]),
            ("D", "RAPPORTS", "#C55A11", "RAPPORTS", [
                "Feuille RAPPORTS : « ＋ Nouveau rapport », puis choisissez la matière, écrivez l'intitulé, choisissez "
                "le type (Individuel ou Groupe) et la date / heure limite.",
                "Cliquez sur la ligne du rapport puis sur « ⇩ Générer les remises attendues » : une ligne par "
                "étudiant (ou par groupe) est créée dans REMISES.",
                "Feuille REMISES : quand quelqu'un dépose son rapport, cliquez sur sa ligne puis sur « ✔ Marquer remis "
                "maintenant » (date et heure notées). Le statut À temps / En retard / Non remis se calcule seul. "
                "Les boutons « Non remis » et « En retard » filtrent la liste.",
            ]),
            ("E", "CONSULTER", VERT, "FICHE ÉTUDIANT", [
                "FICHE ÉTUDIANT : tapez une partie du nom dans « Rechercher » puis Entrée (ou choisissez dans la liste "
                "« Étudiant »). Vous voyez ses présences, absences par matière, cours manqués, rapports et exposés. "
                "Bouton « Fiche en PDF » pour l'imprimer.",
                "RÉCAP ABSENCES : toute la promotion d'un coup d'œil. Rouge = plus de 3 absences dans une matière.",
                "SUIVI ENSEIGNANTS : avancement de chaque matière ; à droite, choisissez un enseignant dans la liste "
                "pour voir toutes ses séances, absences et retards.",
            ]),
            ("F", "IMPRIMER / EXPORTER EN PDF", VERT, "EXPORT PDF", [
                "Feuille EXPORT PDF : choisissez la matière (vide = toutes), le groupe (vide = tous) et "
                "éventuellement une période, puis « Exporter la PRÉSENCE » ou « Exporter le CAHIER DE TEXTE ».",
                "Les PDF sont rangés dans le dossier « Export_PDF » à côté du classeur (un sous-dossier par matière).",
            ]),
            ("G", "BON À SAVOIR", "#7F7F7F", None, [
                "Enregistrez souvent (Ctrl+S). Le bouton « Sauvegarder » de l'accueil crée en plus une copie datée "
                "dans le dossier « Sauvegardes ».",
                "Appuyez sur F9 pour actualiser l'heure et le cours du moment sur l'accueil.",
                "Ne renommez pas les feuilles et ne supprimez pas les tableaux : les macros s'en servent.",
                "Pour modifier une matière, un enseignant ou l'emploi du temps : « Mode administrateur » sur l'accueil.",
            ]),
        ]
        r = 4
        for lettre, titre, coul, cible, etapes in blocs:
            ws.set_row(r, 24)
            ws.write(r, 1, lettre, self.f(bold=True, font_size=13, font_color=BLANC, bg_color=coul, align="center"))
            ws.write(r, 2, "  " + titre, self.f(bold=True, font_size=12, font_color=BLANC, bg_color=coul))
            if cible:
                ws.write_url(r, 4, f"internal:'{cible}'!A1", self.f(bold=True, font_color=BLANC, bg_color=coul,
                                                                    align="center", underline=1),
                             string=f"Ouvrir {cible.lower()} ›")
            else:
                ws.write_blank(r, 4, None, self.f(bg_color=coul))
            ws.write_blank(r, 3, None, self.f(bg_color=coul))
            r += 1
            for i, t in enumerate(etapes, 1):
                lignes = max(1, -(-len(t) // 125))
                ws.set_row(r, 17 * lignes + 6)
                ws.write(r, 1, i, self.f(bold=True, font_color=coul, align="center", valign="top", font_size=12))
                ws.merge_range(r, 2, r, 4, t, self.f(text_wrap=True, valign="top", bottom=4,
                                                     bottom_color="#E7E6E6"))
                r += 1
            r += 1
        legende = r
        ws.write(legende, 2, "Légende des statuts de présence :", self.f(bold=True, font_color=MARINE))
        for i, (st, (fond, txt)) in enumerate(STATUTS.items()):
            ws.write(legende + 1 + i, 1, st, self.f(bold=True, align="center", bg_color=fond, font_color=txt))
            ws.write(legende + 1 + i, 2, SIGNIF[st], self.f(indent=1))
        ws.set_selection("A1")

    # --------------------------------------------------- RÉCAP ABSENCES
    def recap(self):
        nbm = 12
        larg = [5, 32, 10] + [9] * nbm + [8, 8, 9, 9, 11]
        ws = self.feuille("RÉCAP ABSENCES", "shRecap", VERT, "RÉCAPITULATIF DES ABSENCES PAR ÉTUDIANT ET PAR MATIÈRE",
                          "Nombre d'absences (A). Rouge : au-delà du seuil fixé dans PARAMÈTRES ; orange : seuil atteint.",
                          larg, 3 + nbm + 5)
        th = self.f(bold=True, font_color=BLANC, bg_color=BLEU, align="center", text_wrap=True, border=1,
                    border_color=BLANC)
        petit = self.f(font_size=7, font_color="#A6A6A6", align="center")
        ws.write("B6", "N°", th)
        ws.write("C6", "Étudiant", th)
        ws.write("D6", "Groupe", th)
        for j in range(nbm):
            ws.write_formula(4, 4 + j, f'=IFERROR(INDEX(tblMatieres[Code],{j + 1}),"")', petit)
            ws.write_formula(5, 4 + j, f'=IFERROR(INDEX(tblMatieres[Abrégé],{j + 1}),"")', th)
        fin = 4 + nbm
        for c, t in enumerate(["Total A", "Total R", "Taux présence", "Max / matière", "Alerte"]):
            ws.write(5, fin + c, t, self.f(bold=True, font_color=BLANC, bg_color=MARINE, align="center",
                                           text_wrap=True, border=1, border_color=BLANC))
        ws.set_row(5, 40)
        cE, cL = xl_col_to_name(4), xl_col_to_name(fin - 1)
        for k in range(200):
            r = 6 + k
            x = r + 1
            fond = PALE if k % 2 else BLANC
            cel = self.f(bg_color=fond, align="center")
            ws.write_formula(r, 1, f'=IF(C{x}="","",{k + 1})', self.f(bg_color=fond, align="center",
                                                                     font_color=GRIS_TXT))
            ws.write_formula(r, 2, f'=IFERROR(INDEX(tblEtudiants[Nom et prénoms],{k + 1})&"","")',
                             self.f(bg_color=fond, indent=1))
            ws.write_formula(r, 3, f'=IF(C{x}="","",INDEX(tblEtudiants[Groupe],{k + 1})&"")', cel)
            for j in range(nbm):
                col = xl_col_to_name(4 + j)
                ws.write_formula(r, 4 + j, f'=IF(OR($C{x}="",{col}$5=""),"",COUNTIFS(tblPresences[Étudiant],$C{x},tblPresences[Code matière],{col}$5,tblPresences[Statut],"A"))', cel)
            ws.write_formula(r, fin, f'=IF(C{x}="","",SUM({cE}{x}:{cL}{x}))', self.f(bg_color=fond, align="center",
                                                                                   bold=True))
            ws.write_formula(r, fin + 1, f'=IF(C{x}="","",COUNTIFS(tblPresences[Étudiant],C{x},tblPresences[Statut],"R"))', cel)
            ws.write_formula(r, fin + 2, f'=IF(C{x}="","",IFERROR((COUNTIFS(tblPresences[Étudiant],C{x},tblPresences[Statut],"P")+COUNTIFS(tblPresences[Étudiant],C{x},tblPresences[Statut],"R"))/COUNTIF(tblPresences[Étudiant],C{x}),""))',
                             self.f(bg_color=fond, align="center", num_format="0%"))
            ws.write_formula(r, fin + 3, f'=IF(C{x}="","",MAX({cE}{x}:{cL}{x}))', cel)
            ws.write_formula(r, fin + 4, f'=IF(C{x}="","",IF({xl_col_to_name(fin + 3)}{x}>pSeuil,"⚠ ALERTE",""))',
                             self.f(bg_color=fond, align="center", bold=True, font_color=ROUGE))
        plage = f"E7:{cL}206"
        ws.conditional_format(plage, {"type": "formula", "criteria": "=AND(ISNUMBER(E7),E7>pSeuil)",
                                      "format": self.f(bg_color=ROUGE, font_color=BLANC, bold=True)})
        ws.conditional_format(plage, {"type": "formula", "criteria": "=AND(ISNUMBER(E7),E7=pSeuil)",
                                      "format": self.f(bg_color="#F4B183", font_color="#843C0C", bold=True)})
        ws.conditional_format(plage, {"type": "cell", "criteria": "==", "value": 0,
                                      "format": self.f(font_color="#D9D9D9")})
        ws.conditional_format(f"C7:C206", {"type": "formula",
                                           "criteria": f"=AND(ISNUMBER(${xl_col_to_name(fin + 3)}7),${xl_col_to_name(fin + 3)}7>pSeuil)",
                                           "format": self.f(font_color=ROUGE, bold=True)})
        ws.freeze_panes(6, 4)
        ws.autofilter(5, 1, 205, fin + 4)

    # -------------------------------------------------------- EXPORT PDF
    def export(self):
        larg = [26, 30, 50, 12, 12, 12]
        ws = self.feuille("EXPORT PDF", "shExport", VERT, "EXPORT PDF  —  PRÉSENCES PAR MATIÈRE ET PAR GROUPE",
                          "Choisissez les options puis cliquez sur un bouton. Les fichiers sont créés dans le dossier "
                          "« Export_PDF » à côté du classeur (un sous-dossier par matière).", larg, 6)
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        info = self.f(font_color=GRIS_TXT, italic=True, indent=1, text_wrap=True)
        self.section(ws, 4, 1, 6, "OPTIONS")
        lignes = [
            (5, "Matière", "=LstMatieres", "Vide = toutes les matières (un fichier par matière)."),
            (6, "Groupe", "=LstGroupes", "Vide = tous les groupes de la matière."),
            (7, "Mode (si groupe vide)", "=LstModeExport", "Un PDF par groupe, ou une seule liste avec tous les groupes."),
            (8, "Du (facultatif)", None, "Date de début de la période à exporter."),
            (9, "Au (facultatif)", None, "Date de fin de la période à exporter."),
        ]
        for r, t, src, aide in lignes:
            ws.write(r, 1, t, lab)
            nf = "dd/mm/yyyy" if src is None else None
            ws.write_blank(r, 2, None, self.saisie(bold=True, indent=1, num_format=nf) if nf else
                           self.saisie(bold=True, indent=1))
            ws.write(r, 3, aide, info)
            ws.set_row(r, 22)
            if src:
                self.dv_liste(ws, rc(r, 2), src)
            else:
                ws.data_validation(r, 2, r, 2, {"validate": "date", "criteria": ">", "value": dt.date(2020, 1, 1)})
        ws.write("C8", "Un PDF par groupe", self.saisie(bold=True, indent=1))
        ws.set_row(11, 46)
        self.rangee(ws, 11, [
            ("🖨  Exporter la PRÉSENCE (PDF)", dict(macro="ExporterPresencePDF", couleur=VERT, larg=240)),
            ("📖  Exporter le CAHIER DE TEXTE (PDF)", dict(macro="ExporterCahierPDF", couleur=BLEU, larg=270)),
            ("📂  Ouvrir le dossier des PDF", dict(macro="OuvrirDossierPDF", couleur="#7F7F7F", larg=210)),
        ], haut=40, dy=3)
        self.section(ws, 14, 1, 6, "À SAVOIR", MARINE)
        notes = [
            "• Feuille de présence : une colonne par séance (date et heure), P / A / R / E en couleur, totaux et taux "
            "par étudiant, nombre de présents par séance, zone de signature.",
            "• Les colonnes grisées signalent une séance où l'enseignant était absent.",
            "• Un étudiant au-delà du seuil d'absences apparaît en rouge.",
            "• Le cahier de texte PDF reprend pour chaque séance : date, horaire, durée, présence de l'enseignant, type "
            "de séance, contenu et travaux demandés, avec le total d'heures effectuées.",
            "• La fiche individuelle d'un étudiant s'exporte depuis la feuille FICHE ÉTUDIANT.",
        ]
        for i, t in enumerate(notes):
            ws.merge_range(15 + i, 1, 15 + i, 6, t, self.f(text_wrap=True, indent=1))
            ws.set_row(15 + i, 30)
        ws.write("B22", "Dossier des exports :", lab)
        ws.merge_range("C22:G22", "", self.f(font_color=BLEU, indent=1))
        ws.write_formula("C22", '=IFERROR(LEFT(CELL("filename",A1),FIND("[",CELL("filename",A1))-1)&"Export_PDF","(enregistrez d\'abord le classeur)")',
                         self.f(font_color=BLEU, indent=1))

    # -------------------------------------------------------- ÉTUDIANTS
    def etudiants_f(self):
        larg = [6, 34, 11, 12, 32, 15, 3, 12, 10]
        ws = self.feuille("ÉTUDIANTS", "shEtudiants", "#7F7F7F", "LISTE DES ÉTUDIANTS",
                          "Le « Groupe » indiqué ici est le groupe par défaut (= groupes de TPS, Travaux Pratiques Spécialisés). "
                          "Il sert pour toutes les matières sauf celles qui ont leurs propres groupes (feuille GROUPES).", larg, 9)
        ws.set_row(3, 30)
        self.rangee(ws, 3, [
            ("＋  Ajouter un étudiant", dict(macro="AjouterLigne", couleur="#7F7F7F", larg=170)),
            ("✖  Supprimer la sélection", dict(macro="SupprimerLignes", couleur=ROUGE, larg=190)),
        ])
        colonnes = [
            {"header": "N°", "format": self.f(locked=False, align="center")},
            {"header": "Nom et prénoms", "format": self.f(locked=False, bold=True)},
            {"header": "Groupe", "format": self.f(locked=False, align="center")},
            {"header": "Rôle", "format": self.f(locked=False, align="center")},
            {"header": "Email", "format": self.f(locked=False, font_color=BLEU)},
            {"header": "Contact", "format": self.f(locked=False, align="center", num_format="00 00 00 00 00")},
        ]
        self.table(ws, 5, 1, "tblEtudiants", colonnes, self.etudiants or None, style="Table Style Medium 2")
        self.dv_liste(ws, f"D7:D{NB_LIGNES_TABLE}", "=LstGroupes")
        self.section(ws, 5, 8, 9, "EFFECTIFS")
        for k in range(15):
            r = 6 + k
            ws.write_formula(r, 8, f'=IFERROR(INDEX(tblListeGroupes[Groupe],{k + 1})&"","")',
                             self.f(indent=1, bg_color=PALE if k % 2 else BLANC))
            ws.write_formula(r, 9, f'=IF(I{r + 1}="","",COUNTIF(tblEtudiants[Groupe],I{r + 1}))',
                             self.f(align="center", bold=True, bg_color=PALE if k % 2 else BLANC))
        ws.write("I22", "Total", self.f(bold=True, indent=1, top=1))
        ws.write_formula("J22", "=COUNTA(tblEtudiants[Nom et prénoms])", self.f(bold=True, align="center", top=1))
        ws.freeze_panes(6, 0)

    # ---------------------------------------------------------- GROUPES
    def groupes(self):
        larg = [12, 34, 11, 12, 3] + [21] * 16
        ws = self.feuille("GROUPES", "shGroupes", "#7F7F7F", "GROUPES PAR MATIÈRE",
                          "Par défaut, chaque matière utilise les groupes de la feuille ÉTUDIANTS. Si une matière a des "
                          "groupes différents : choisissez-la, cliquez sur « Préparer », puis modifiez la colonne Groupe "
                          "(ou supprimez les étudiants qui ne suivent pas la matière).", larg, 15)
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        ws.write("B5", "Matière", lab)
        ws.write_blank("C5", None, self.saisie(bold=True, indent=1))
        self.dv_liste(ws, "C5", "=LstMatieres")
        ws.merge_range("B6:E6", "", self.f(bold=True, font_color=OR, indent=1))
        ws.write_formula("B6", '=IF(C5="","Choisissez une matière.",IF(COUNTIF(tblGroupes[Code matière],LEFT(C5,FIND(" ",C5&" ")-1))>0,"► Groupes SPÉCIFIQUES à "&C5&" (tableau ci-dessous)","► "&C5&" utilise les groupes PAR DÉFAUT (feuille ÉTUDIANTS)"))',
                         self.f(bold=True, font_color=OR, indent=1))
        ws.set_row(4, 30)
        self.rangee(ws, 4, [
            ("⇩  Préparer les groupes de cette matière", dict(macro="PreparerGroupesMatiere", couleur=ORANGE,
                                                             larg=270)),
            ("↺  Revenir aux groupes par défaut", dict(macro="SupprimerGroupesMatiere", couleur="#7F7F7F", larg=240)),
            ("＋ Ligne", dict(macro="AjouterLigne", couleur="#7F7F7F", larg=80)),
            ("✖ Supprimer", dict(macro="SupprimerLignes", couleur=ROUGE, larg=105)),
        ], x0=440)
        colonnes = [{"header": "Code matière", "format": self.f(locked=False, align="center")},
                    {"header": "Étudiant", "format": self.f(locked=False)},
                    {"header": "Groupe", "format": self.f(locked=False, align="center", bold=True)},
                    {"header": "Rôle", "format": self.f(locked=False, align="center")}]
        self.table(ws, 8, 1, "tblGroupes", colonnes, style="Table Style Medium 7")
        self.dv_liste(ws, f"B10:B{NB_LIGNES_TABLE}", "=LstMatieres")
        self.dv_liste(ws, f"C10:C{NB_LIGNES_TABLE}", "=LstEtudiants")
        self.dv_liste(ws, f"D10:D{NB_LIGNES_TABLE}", "=LstGroupes")
        # Composition
        self.section(ws, 7, 6, 21, "COMPOSITION DES GROUPES POUR LA MATIÈRE CHOISIE (mise à jour automatique)", VERT)
        th = self.f(bold=True, font_color=BLANC, bg_color=VERT, align="center", text_wrap=True)
        for k in range(16):
            col = 6 + k
            ws.write_blank(8, col, None, th)
            for r in range(9, 60):
                ws.write_blank(r, col, None, self.f(font_size=9, indent=1, bottom=4, bottom_color="#E7E6E6"))
        ws.set_row(8, 30)

    # --------------------------------------------------------- MATIÈRES
    def matieres(self):
        larg = [12, 10, 46, 15, 24, 9, 9, 9, 11, 10, 24, 3, 24, 16, 28]
        ws = self.feuille("MATIÈRES", "shMatieres", "#7F7F7F", "MATIÈRES ET ENSEIGNANTS  —  SEMESTRE 9",
                          "D'après l'emploi du temps officiel. Le volume prévu sert au calcul de l'avancement. "
                          "Modifications : mode administrateur (mot de passe).", larg, 10)
        colonnes = [
            {"header": "Code", "format": self.f(bold=True, align="center"), "total_string": "TOTAL"},
            {"header": "UE", "format": self.f(align="center")},
            {"header": "Intitulé", "format": self.f(text_wrap=True)},
            {"header": "Abrégé", "format": self.f(align="center")},
            {"header": "Enseignant(s)", "format": self.f()},
            {"header": "Cours (h)", "format": self.f(align="center"),
             "total_function": "=SUBTOTAL(109,tblMatieres[Cours (h)])"},
            {"header": "TD (h)", "format": self.f(align="center"),
             "total_function": "=SUBTOTAL(109,tblMatieres[TD (h)])"},
            {"header": "TP (h)", "format": self.f(align="center"),
             "total_function": "=SUBTOTAL(109,tblMatieres[TP (h)])"},
            {"header": "Volume total (h)", "format": self.f(align="center", bold=True),
             "formula": "=[@[Cours (h)]]+[@[TD (h)]]+[@[TP (h)]]",
             "total_function": "=SUBTOTAL(109,tblMatieres[Volume total (h)])"},
            {"header": "Couleur", "format": self.f(align="center", font_size=8)},
            {"header": "Libellé", "format": self.f(font_size=9, font_color=GRIS_TXT),
             "formula": '=[@Code]&" – "&[@Abrégé]'},
        ]
        donnees = [[m[0], m[1], m[2], m[3], m[4], m[5], m[6], m[7], None, m[8], None] for m in MATIERES]
        self.table(ws, 5, 1, "tblMatieres", colonnes, donnees, style="Table Style Medium 2", total=True)
        for i, m in enumerate(MATIERES):
            ws.write(6 + i, 10, m[8], self.f(align="center", font_size=8, bg_color=m[8], font_color=GRIS_TXT))
            ws.set_row(6 + i, 30)
        self.dv_liste(ws, "F7:F60", "=LstEnseignants", libre=True)
        # Enseignants
        self.section(ws, 3, 13, 15, "ENSEIGNANTS (liste déroulante du classeur)", VERT)
        self.table(ws, 5, 13, "tblEnseignants", [
            {"header": "Enseignant", "format": self.f(locked=False, bold=True)},
            {"header": "Téléphone", "format": self.f(locked=False, align="center")},
            {"header": "Email", "format": self.f(locked=False, font_color=BLEU)}],
            [[e, None, None] for e in ENSEIGNANTS], style="Table Style Medium 7")
        self.bouton(ws, 4, 13, "＋ Ajouter un enseignant", macro="AjouterEnseignant", couleur=VERT, larg=170,
                    haut=20, dy=1)

    # ----------------------------------------------------- EMPLOI DU TEMPS
    def edt(self):
        larg = [10, 7, 8, 8, 11, 30, 20, 9, 8, 9, 3, 11, 19, 19, 19, 19, 19, 19]
        ws = self.feuille("EMPLOI DU TEMPS", "shEDT", "#7F7F7F", "EMPLOI DU TEMPS  —  du 28/09/2026 au 22/01/2027",
                          "Le tableau de gauche pilote « DÉMARRER LE COURS » et l'accueil (cours du moment, prochain cours). "
                          "La grille de droite est une vue de l'emploi du temps officiel.", larg, 18)
        colonnes = [
            {"header": "Jour", "format": self.f(bold=True, align="center")},
            {"header": "N° jour", "format": self.f(align="center", font_color=GRIS_TXT),
             "formula": '=IFERROR(MATCH([@Jour],{"Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"},0),"")'},
            {"header": "Début", "format": self.f(align="center", num_format="hh:mm")},
            {"header": "Fin", "format": self.f(align="center", num_format="hh:mm")},
            {"header": "Code matière", "format": self.f(align="center", bold=True)},
            {"header": "Matière", "format": self.f(),
             "formula": '=IFERROR(XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Intitulé]),"")'},
            {"header": "Enseignant", "format": self.f(),
             "formula": '=IFERROR(XLOOKUP([@[Code matière]],tblMatieres[Code],tblMatieres[Enseignant(s)]),"")'},
            {"header": "Salle", "format": self.f(align="center")},
            {"header": "Durée (h)", "format": self.f(align="center", num_format="0.0"),
             "formula": '=IF(OR([@Début]="",[@Fin]=""),"",ROUND(([@Fin]-[@Début])*24,2))'},
            {"header": "Minutes avant", "format": self.f(align="center", font_color="#A6A6A6", font_size=8),
             "formula": '=IF(OR([@[N° jour]]="",[@Début]=""),"",MOD(([@[N° jour]]-1)*1440+ROUND([@Début]*1440,0)-((WEEKDAY(TODAY(),2)-1)*1440+ROUND(MOD(NOW(),1)*1440,0)),10080))'},
        ]
        donnees = [[j, None, heure(d), heure(f), code, None, None, salle, None, None] for j, d, f, code, salle in EDT]
        self.table(ws, 5, 1, "tblEDT", colonnes, donnees, style="Table Style Medium 2")
        self.dv_liste(ws, f"B7:B200", "=LstJours")
        self.dv_liste(ws, f"F7:F200", "=LstMatieres")
        # Grille
        g0 = 12
        th = self.f(bold=True, font_color=BLANC, bg_color=MARINE, align="center", border=1, border_color=BLANC)
        ws.write(5, g0, "Horaire", th)
        for j, jour in enumerate(JOURS):
            ws.write(5, g0 + 1 + j, jour.upper(), th)
        for h in range(8, 19):
            r = 6 + h - 8
            ws.write(r, g0, f"{h:02d}h - {h + 1:02d}h", self.f(bold=True, align="center", bg_color=BLEU_CLAIR,
                                                               font_color=MARINE, border=1, border_color=BLANC))
            ws.set_row(r, 30)
            for j in range(6):
                ws.write_blank(r, g0 + 1 + j, None, self.f(border=1, border_color="#E7E6E6"))
        for jour, d, f, code, salle in EDT:
            j = JOURS.index(jour)
            txt = f"{code}\n{ABREGE[code]}\n{ENSEIGNANT[code]}" + (f" · {salle}" if salle else "")
            fmt = self.f(bg_color=COULEUR[code], align="center", text_wrap=True, font_size=9, bold=True,
                         font_color=MARINE, border=2, border_color=BLANC)
            r1, r2 = 6 + d - 8, 6 + f - 8 - 1
            if r1 == r2:
                ws.write(r1, g0 + 1 + j, txt, fmt)
            else:
                ws.merge_range(r1, g0 + 1 + j, r2, g0 + 1 + j, txt, fmt)
        ws.merge_range(6, g0 + 5, 9, g0 + 5, "TPE\nen salle de cours",
                       self.f(bg_color="#EDEDED", align="center", text_wrap=True, font_size=9, italic=True,
                              font_color=GRIS_TXT, border=2, border_color=BLANC))
        ws.merge_range(11, g0 + 1, 11, g0 + 6, "PAUSE", self.f(bg_color="#F2F2F2", align="center",
                                                                font_color="#A6A6A6", italic=True))

    # ------------------------------------------------------- PARAMÈTRES
    def parametres(self):
        larg = [44, 36, 3, 9, 14, 13, 13, 13, 11, 30, 3, 3, 14]
        ws = self.feuille("PARAMÈTRES", "shParam", "#7F7F7F", "PARAMÈTRES ET LISTES",
                          "Les valeurs en jaune sont modifiables. Les listes servent aux menus déroulants de tout le classeur.",
                          larg, 13)
        self.ws_param = ws
        self.section(ws, 4, 1, 2, "PARAMÈTRES GÉNÉRAUX")
        lab = self.f(bold=True, font_color=MARINE, bg_color=PALE, indent=1, border=1, border_color=BORDURE)
        self.cell_param = {}
        for i, (nom, lib, val) in enumerate(PARAMS):
            r = 5 + i
            ws.write(r, 1, lib, lab)
            val = getattr(self, "params_reprise", {}).get(nom, val)
            if val is None:
                val = ""
            nf = "dd/mm/yyyy" if isinstance(val, dt.date) else None
            fmt = self.saisie(bold=True, indent=1, num_format=nf) if nf else self.saisie(bold=True, indent=1)
            if isinstance(val, dt.date):
                ws.write_datetime(r, 2, dt.datetime.combine(val, dt.time()), fmt)
            else:
                ws.write(r, 2, val, fmt)
            ws.set_row(r, 21)
            self.cell_param[nom] = f"'PARAMÈTRES'!$C${r + 1}"
        # Listes
        self.section(ws, 4, 4, 10, "LISTES DÉROULANTES")
        th = self.f(bold=True, font_color=MARINE, bg_color=BLEU_CLAIR, align="center", bottom=1, text_wrap=True)
        self.plages_listes = {}
        col = 4
        for nom, (titre, valeurs) in LISTES.items():
            ws.write(5, col, titre, th)
            for i, v in enumerate(valeurs):
                if nom == "LstStatuts":
                    fond, txt = STATUTS[v]
                    ws.write(6 + i, col, v, self.f(align="center", bold=True, bg_color=fond, font_color=txt))
                else:
                    ws.write(6 + i, col, v, self.f(align="center"))
            lettre = xl_col_to_name(col)
            self.plages_listes[nom] = f"='PARAMÈTRES'!${lettre}$7:${lettre}${6 + len(valeurs)}"
            col += 1
        ws.set_row(5, 30)
        # Groupes
        self.section(ws, 4, 13, 13, "GROUPES", VERT)
        self.table(ws, 5, 13, "tblListeGroupes", [{"header": "Groupe", "format": self.f(locked=False,
                                                                                     align="center")}],
                   [[f"Groupe {i}"] for i in range(1, 8)], style="Table Style Light 14")
        self.bouton(ws, 13, 13, "＋ Groupe", macro="AjouterLigne", couleur=VERT, larg=95, haut=24, dy=4)
        # Palette
        self.section(ws, 18, 1, 2, "PALETTE DE COULEURS DU CLASSEUR")
        palette = [(MARINE, "Bleu marine — titres, en-têtes principaux"), (BLEU, "Bleu — en-têtes de tableaux"),
                   (BLEU_CLAIR, "Bleu clair — sous-en-têtes"), (PALE, "Bleu pâle — lignes alternées, cartes"),
                   (ORANGE, "Orange — saisie (cours du jour, cahier, présences)"),
                   (VERT, "Vert — consultation (fiches, suivi, exports)"), (ROUGE, "Rouge — alertes, suppressions"),
                   (OR, "Or — prochain cours, exposés"), ("#FFFDF2", "Jaune clair — cellules à remplir")]
        for i, (c, t) in enumerate(palette):
            ws.write(19 + i, 1, t, self.f(indent=1))
            ws.write(19 + i, 2, c, self.f(bg_color=c, align="center", font_color=BLANC if c in (MARINE, BLEU, ORANGE,
                                                                                               VERT, ROUGE, OR)
                                          else MARINE, bold=True))
        for i, (s, (fond, txt)) in enumerate(STATUTS.items()):
            ws.write(28 + i, 1, f"{s} — {SIGNIF[s]}", self.f(indent=1))
            ws.write(28 + i, 2, f"{fond} / {txt}", self.f(bg_color=fond, font_color=txt, bold=True, align="center"))
        ws.write("B34", f"Mot de passe des feuilles protégées : {MDP}  (modifiable dans le module VBA modOutils).",
                 self.f(italic=True, font_color=GRIS_TXT))

    # -------------------------------------------------------- IMPRESSION
    def impression(self):
        ws = self.wb.add_worksheet("IMPRESSION")
        ws.set_vba_name("shImpression")
        ws.hide()
        ws.write("A1", "Feuille de travail des exports PDF (masquée).")

    # ------------------------------------------------------------- noms
    def noms(self):
        for nom, ref in self.cell_param.items():
            self.wb.define_name(nom, "=" + ref)
        for nom, ref in self.plages_listes.items():
            self.wb.define_name(nom, ref)
        self.wb.define_name("LstCodes", "=tblMatieres[Code]")
        self.wb.define_name("LstMatieres", "=tblMatieres[Libellé]")
        self.wb.define_name("LstEnseignants", "=tblEnseignants[Enseignant]")
        self.wb.define_name("LstEtudiants", "=tblEtudiants[Nom et prénoms]")
        self.wb.define_name("LstGroupes", "=tblListeGroupes[Groupe]")
        self.wb.define_name("LstSeances", "=tblSeances[ID]")
        self.wb.define_name("LstRapports", "=tblRapports[ID]")

    # -------------------------------------------------------------- VBA
    def vba(self):
        modules = []
        noms_feuilles = [ws.vba_codename for ws in self.wb.worksheets()]
        sources = {}
        for fichier in sorted(os.listdir(VBA_DIR)):
            nom, ext = os.path.splitext(fichier)
            with open(os.path.join(VBA_DIR, fichier), encoding="utf-8") as fh:
                sources[nom] = (ext, fh.read())
        modules.append({"name": "ThisWorkbook", "type": "workbook", "code": sources["ThisWorkbook"][1]})
        for n in noms_feuilles:
            code = sources[n][1] if n in sources else "Option Explicit\n"
            modules.append({"name": n, "type": "sheet", "code": code})
        for nom, (ext, code) in sources.items():
            if ext == ".bas":
                modules.append({"name": nom, "type": "std", "code": code})
        chemin = os.path.join(os.path.dirname(self.wb.filename) or ".", "vbaProject.bin")
        build_vba_project(modules, chemin)
        self.wb.add_vba_project(chemin)
        self._vba_tmp = chemin

    # ------------------------------------------------------------ démo
    def demo_seances(self):
        random.seed(4)
        lignes = []
        d = dt.date(2026, 9, 28)
        n = 0
        while d <= dt.date(2026, 10, 6):
            for jour, h1, h2, code, salle in EDT:
                if JOURS.index(jour) == d.weekday():
                    n += 1
                    prof = random.choice(["Présent"] * 6 + ["Absent", "Retard"])
                    lignes.append([f"S{n:04d}", dt.datetime.combine(d, dt.time()), None, heure(h1), heure(h2), None,
                                   code, None, None, prof, "", "Cours", f"Chapitre {n % 5 + 1} : notions de base",
                                   "Lire le chapitre suivant" if n % 3 == 0 else "", None, None, None, None, None, ""])
            d += dt.timedelta(days=1)
        self._demo_seances = lignes
        return lignes

    def demo_presences(self):
        random.seed(7)
        out = []
        for s in getattr(self, "_demo_seances", []):
            if s[9] == "Absent":
                continue
            for e in self.etudiants:
                st = random.choices("PARE", [80, 12, 6, 2])[0]
                out.append([s[0], e[1], e[2], s[6], dict((m[0], m[2]) for m in MATIERES)[s[6]], s[1], s[3], st,
                            None, "", dt.datetime.combine(s[1].date(), s[3])])
        return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sortie")
    ap.add_argument("--etudiants")
    ap.add_argument("--demo", action="store_true")
    ap.add_argument("--apercu", action="store_true", help="une page par feuille (contrôle visuel)")
    ap.add_argument("--reprendre", help="ancien classeur SUIVI_GC5 dont on reprend les saisies")
    a = ap.parse_args()
    etu = lire_etudiants(a.etudiants) if a.etudiants else []
    c = Classeur(a.sortie, etu, a.demo, a.apercu)
    if a.reprendre:
        c.reprise, c.params_reprise = lire_reprise(a.reprendre)
        print("Reprise :", {k: len(v) for k, v in c.reprise.items()})
    c.construire()
    os.remove(c._vba_tmp)
    print(f"Classeur créé : {a.sortie}  ({len(etu)} étudiants)")


if __name__ == "__main__":
    main()
